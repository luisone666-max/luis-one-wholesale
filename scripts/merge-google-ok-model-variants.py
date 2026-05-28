import csv
import hashlib
import json
import math
import re
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
PREVIEW_DIR = ROOT / "tmp-imports" / "google-ok" / "preview"
PRODUCTS_CSV = PREVIEW_DIR / "ok-products-preview.csv"
VARIANTS_CSV = PREVIEW_DIR / "ok-variants-preview.csv"
BRAND_SPLIT_VARIANT_THRESHOLD = 8
BRAND_SPLIT_MIN_RECOGNIZED_VARIANTS = 4
MODEL_SPLIT_MIN_RECOGNIZED_VARIANTS = 6

MODEL_MARKERS = {
    "ADV",
    "ADV150",
    "ADV160",
    "AEROX",
    "AEROX155",
    "BARAKO",
    "BEAT",
    "BEATFI",
    "BURGMAN",
    "C100",
    "CG125",
    "CLICK",
    "CLICK125",
    "CLICK150",
    "CRYPTON",
    "CT100",
    "CT125",
    "CT150",
    "DASH",
    "DASH110",
    "DREAM",
    "DIO",
    "F16",
    "FZ16",
    "FURY",
    "FURY125",
    "GP125",
    "GRAVIS",
    "GY6",
    "HD3",
    "LC135",
    "LC150",
    "MIO",
    "MIO125",
    "NEX",
    "NMAX",
    "NMAX155",
    "NOUVO",
    "PCX",
    "PCX150",
    "PCX160",
    "RAIDER",
    "RAIDER150",
    "ROUSER",
    "RS100",
    "RS125",
    "RS150",
    "RUSI",
    "RUSI125",
    "RUSI150",
    "SHOGUN",
    "SKYDRIVE",
    "SKYGO",
    "SMASH",
    "SMASH110",
    "SMASH115",
    "SNIPER",
    "SNIPER135",
    "SNIPER150",
    "SNIPER155",
    "SOUL",
    "SPORTY",
    "STX",
    "SUPREMO",
    "TMX",
    "TMX125",
    "TMX155",
    "VEGA",
    "WAVE",
    "WAVE100",
    "WAVE100R",
    "WAVE110",
    "WAVE125",
    "WAVE125I",
    "WAVE125R",
    "WIND125",
    "XLR200",
    "XR200",
    "XRM",
    "XRM110",
    "XRM125",
    "X1",
    "X1R",
    "YTX",
}

MODEL_TOKEN_BRANDS = {
    "ADV": "Honda",
    "ADV150": "Honda",
    "ADV160": "Honda",
    "BEAT": "Honda",
    "BEATFI": "Honda",
    "C100": "Honda",
    "CB110": "Honda",
    "CB125": "Honda",
    "CG150": "Honda",
    "CG125": "Honda",
    "CLICK": "Honda",
    "CLICK125": "Honda",
    "CLICK150": "Honda",
    "CLICK160": "Honda",
    "DASH": "Honda",
    "DASH110": "Honda",
    "DIO": "Honda",
    "DREAM": "Honda",
    "PCX": "Honda",
    "PCX150": "Honda",
    "PCX160": "Honda",
    "RS100": "Honda",
    "RS125": "Honda",
    "RS150": "Honda",
    "SCOOPY": "Honda",
    "SCOOPYFI": "Honda",
    "SONIC": "Honda",
    "SONIC150": "Honda",
    "STX": "Honda",
    "STX125": "Honda",
    "SUPREMO": "Honda",
    "TIGER": "Honda",
    "TMX": "Honda",
    "TMX125": "Honda",
    "TMX155": "Honda",
    "WAVE": "Honda",
    "WAVE100": "Honda",
    "WAVE100R": "Honda",
    "WAVE110": "Honda",
    "WAVE125": "Honda",
    "WAVE125I": "Honda",
    "WAVE125R": "Honda",
    "XLR200": "Honda",
    "XR200": "Honda",
    "XRM": "Honda",
    "XRM110": "Honda",
    "XRM125": "Honda",
    "AEROX": "Yamaha",
    "AEROX155": "Yamaha",
    "CRYPTON": "Yamaha",
    "F16": "Yamaha",
    "FZ16": "Yamaha",
    "GRAVIS": "Yamaha",
    "LC135": "Yamaha",
    "LC150": "Yamaha",
    "MIO": "Yamaha",
    "MIO125": "Yamaha",
    "NMAX": "Yamaha",
    "NMAX155": "Yamaha",
    "NOUVO": "Yamaha",
    "SNIPER": "Yamaha",
    "SNIPER135": "Yamaha",
    "SNIPER150": "Yamaha",
    "SNIPER155": "Yamaha",
    "SOUL": "Yamaha",
    "VEGA": "Yamaha",
    "X1": "Yamaha",
    "X1R": "Yamaha",
    "YTX": "Yamaha",
    "YTX125": "Yamaha",
    "AXELO": "Suzuki",
    "BURGMAN": "Suzuki",
    "GD110": "Suzuki",
    "GP125": "Suzuki",
    "GS125": "Suzuki",
    "NEX": "Suzuki",
    "RAIDER": "Suzuki",
    "RAIDER150": "Suzuki",
    "SHOGUN": "Suzuki",
    "SHOGUN110": "Suzuki",
    "SKYDRIVE": "Suzuki",
    "SMASH": "Suzuki",
    "SMASH110": "Suzuki",
    "SMASH115": "Suzuki",
    "BAJAJ": "Kawasaki/Bajaj",
    "BARAKO": "Kawasaki/Bajaj",
    "BC175": "Kawasaki/Bajaj",
    "CT100": "Kawasaki/Bajaj",
    "CT125": "Kawasaki/Bajaj",
    "CT150": "Kawasaki/Bajaj",
    "FURY": "Kawasaki/Bajaj",
    "FURY125": "Kawasaki/Bajaj",
    "HD3": "Kawasaki/Bajaj",
    "HDIII": "Kawasaki/Bajaj",
    "ROUSER": "Kawasaki/Bajaj",
    "WIND125": "Kawasaki/Bajaj",
    "RUSI": "Rusi",
    "RUSI125": "Rusi",
    "RUSI150": "Rusi",
    "GY6": "GY6 Scooter",
    "LIFAN": "Lifan",
    "LIFAN110": "Lifan",
    "LIFAN125": "Lifan",
    "LIFAN150": "Lifan",
}

MODEL_TOKEN_FAMILIES = {
    "ADV": "ADV",
    "ADV150": "ADV",
    "ADV160": "ADV",
    "BEAT": "Beat",
    "BEATFI": "Beat",
    "C100": "Dream/C100",
    "CB110": "CB",
    "CB125": "CB",
    "CG125": "TMX/CG",
    "CG150": "TMX/CG",
    "CLICK": "Click",
    "CLICK125": "Click",
    "CLICK150": "Click",
    "CLICK160": "Click",
    "DASH": "Dash",
    "DASH110": "Dash",
    "DIO": "Dio",
    "DREAM": "Dream/C100",
    "PCX": "PCX",
    "PCX150": "PCX",
    "PCX160": "PCX",
    "RS100": "RS",
    "RS125": "RS",
    "RS150": "RS",
    "SCOOPY": "Scoopy",
    "SCOOPYFI": "Scoopy",
    "SONIC": "Sonic",
    "SONIC150": "Sonic",
    "STX": "STX/Supremo",
    "STX125": "STX/Supremo",
    "SUPREMO": "STX/Supremo",
    "TIGER": "Tiger",
    "TMX": "TMX/CG",
    "TMX125": "TMX/CG",
    "TMX155": "TMX/CG",
    "WAVE": "Wave",
    "WAVE100": "Wave",
    "WAVE100R": "Wave",
    "WAVE110": "Wave",
    "WAVE125": "Wave",
    "WAVE125I": "Wave",
    "WAVE125R": "Wave",
    "XLR200": "XR/XLR",
    "XR200": "XR/XLR",
    "XRM": "XRM",
    "XRM110": "XRM",
    "XRM125": "XRM",
    "AEROX": "Aerox",
    "AEROX155": "Aerox",
    "CRYPTON": "Crypton",
    "F16": "FZ16",
    "FZ16": "FZ16",
    "GRAVIS": "Gravis",
    "LC135": "LC",
    "LC150": "LC",
    "MIO": "Mio",
    "MIO125": "Mio",
    "NMAX": "NMAX",
    "NMAX155": "NMAX",
    "NOUVO": "Nouvo",
    "SNIPER": "Sniper",
    "SNIPER135": "Sniper",
    "SNIPER150": "Sniper",
    "SNIPER155": "Sniper",
    "SOUL": "Soul",
    "VEGA": "Vega",
    "X1": "X1/X1R",
    "X1R": "X1/X1R",
    "YTX": "YTX",
    "YTX125": "YTX",
    "AXELO": "Axelo",
    "BURGMAN": "Burgman",
    "GD110": "GD/GS",
    "GP125": "GD/GS",
    "GS125": "GD/GS",
    "NEX": "NEX",
    "RAIDER": "Raider",
    "RAIDER150": "Raider",
    "SHOGUN": "Shogun",
    "SHOGUN110": "Shogun",
    "SKYDRIVE": "Skydrive",
    "SMASH": "Smash",
    "SMASH110": "Smash",
    "SMASH115": "Smash",
    "BAJAJ": "Bajaj",
    "BARAKO": "Barako",
    "BC175": "Barako",
    "CT100": "CT",
    "CT125": "CT",
    "CT150": "CT",
    "FURY": "Fury",
    "FURY125": "Fury",
    "HD3": "HD3",
    "HDIII": "HD3",
    "ROUSER": "Rouser",
    "WIND125": "Wind",
    "RUSI": "Rusi",
    "RUSI125": "Rusi",
    "RUSI150": "Rusi",
    "GY6": "GY6",
    "LIFAN": "Lifan",
    "LIFAN110": "Lifan",
    "LIFAN125": "Lifan",
    "LIFAN150": "Lifan",
}

MODEL_MARKERS.update(MODEL_TOKEN_BRANDS.keys())
MODEL_MARKERS.update(MODEL_TOKEN_FAMILIES.keys())

KNOWN_BASE_PREFIXES = [
    "AIR CLEANER TUBE",
    "AIR FILTER ELEMENT",
    "BATTERY COVER",
    "BODY COVER",
    "BRAKE ARM",
    "BRAKE CAM",
    "BRAKE LEVER",
    "BRAKE ROD",
    "BRAKE SHOE",
    "CALIPER SET",
    "CENTER COVER",
    "CENTER CONSOLE",
    "CLUTCH LEVER",
    "CONSOLE",
    "CYLINDER BLOCK",
    "CYLINDER",
    "DISC COVER",
    "DISC PLATE",
    "ENGINE SPROCKET",
    "FAN COVER",
    "FLERINGS CENTER COVER",
    "FRONT COWLING",
    "FRONT COVER",
    "FRONT FORK OIL SEAL",
    "FRONT PANEL CENTER COVER",
    "FRONT PANEL COVER",
    "FRONT SHOCK",
    "FRONT TOP COVER",
    "FUEL CAP",
    "FUEL COCK",
    "FUEL FILTER",
    "FUEL HOSE",
    "FUEL PUMP FILTER",
    "FUEL PUMP MOTOR",
    "HANDLE COVER",
    "HANDLE SWITCH LH",
    "HANDLE SWITCH RH",
    "HEAD LIGHT COWLING",
    "HEAD LIGHT LENS",
    "HEADLIGHT LENS",
    "IGNITION COVER",
    "INNER LEG SHIELD",
    "LEG SHIELD COVER",
    "LEG SHIELD",
    "LEGSHIELD COVER",
    "LONG SIDE COVER",
    "MAIN SWITCH SET",
    "MAIN SWITCH",
    "MATTING RUBBER",
    "MUGS SET",
    "OIL BREATHER HOSE",
    "OIL FILTER",
    "OIL SEAL KIT",
    "OIL SEAL SET",
    "OIL SEAL",
    "PISTON KIT",
    "PISTON RING",
    "REAR COWLING",
    "REAR HUB",
    "REGULATOR",
    "SEAT ASSY",
    "SHIFTER",
    "SHOCK",
    "SIDE COVER",
    "SIDE MIRROR",
    "SIGNAL LIGHT COVER",
    "SPROCKET BLACK",
    "SPROCKET",
    "STATOR COIL",
    "TAIL LIGHT LED",
    "TAIL LIGHT",
    "TIRE TUBE TYPE",
    "TIRE TUBELESS",
    "TOP GASKET",
]


def clean(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip())


def slugify(value, fallback="item"):
    text = clean(value).lower()
    text = re.sub(r"[^a-z0-9]+", "-", text)
    return re.sub(r"-+", "-", text).strip("-") or fallback


def stable_id(*parts):
    raw = "|".join(clean(part) for part in parts)
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:8]


def read_csv(path):
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def write_csv(path, rows, fieldnames):
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def strip_supplier_prefix(title, sheet):
    text = clean(title).upper()
    for prefix in ["OK"]:
        if text.startswith(prefix + " "):
            text = text[len(prefix) + 1 :]
    return text


def normalize_token(token):
    return re.sub(r"[^A-Z0-9]", "", token.upper())


def tokens_for_brand_detection(*values):
    tokens = []
    for value in values:
        for token in re.split(r"[^A-Z0-9]+", clean(value).upper()):
            normalized = normalize_token(token)
            if normalized:
                tokens.append(normalized)
    return tokens


def has_model_marker(text):
    tokens = tokens_for_brand_detection(text)
    return any(token in MODEL_MARKERS for token in tokens)


def motorcycle_brand_for_variant(variant):
    tokens = tokens_for_brand_detection(variant.get("model", ""), variant.get("variant_name", ""), variant.get("size", ""))
    brands = []
    for token in tokens:
        brand = MODEL_TOKEN_BRANDS.get(token)
        if brand and brand not in brands:
            brands.append(brand)
    if not brands:
        return "Other Fitments"
    if len(brands) == 1:
        return brands[0]
    if "Kawasaki/Bajaj" in brands:
        return "Kawasaki/Bajaj"
    return brands[0]


def motorcycle_model_family_for_variant(variant):
    tokens = tokens_for_brand_detection(variant.get("model", ""), variant.get("variant_name", ""), variant.get("size", ""))
    for token in tokens:
        family = MODEL_TOKEN_FAMILIES.get(token)
        if family:
            return family
    return "Other Models"


def model_base(title, sheet):
    text = strip_supplier_prefix(title, sheet)
    text = re.sub(r"\s+", " ", text).strip()
    brand_prefix = ""
    match_text = text
    for prefix in ["MSM", "HACHI"]:
        if text.startswith(prefix + " "):
            brand_prefix = prefix
            match_text = text[len(prefix) + 1 :].strip()
            break
    for prefix in sorted(KNOWN_BASE_PREFIXES, key=len, reverse=True):
        if match_text == prefix:
            return text, ""
        if match_text.startswith(prefix + " "):
            tail = match_text[len(prefix) + 1 :].strip()
            if has_model_marker(tail):
                base = f"{brand_prefix} {prefix}".strip()
                return base, tail
    tokens = match_text.split()
    for index, token in enumerate(tokens):
        if normalize_token(token) in MODEL_MARKERS and index >= 2:
            base = " ".join(tokens[:index])
            return f"{brand_prefix} {base}".strip(), " ".join(tokens[index:])
    return text, ""


def should_merge(group, product_meta):
    if len(group) < 2:
        return False
    tails = [clean(product_meta[product["preview_id"]]["tail"]) for product in group]
    if any(has_model_marker(tail) for tail in tails):
        return True
    return any(not clean(product.get("main_image_url")) for product in group)


def merged_product_name(base):
    return "OK " + clean(base).title().replace("Led", "LED").replace("Cdi", "CDI")


def append_tail_to_variant(tail, variant_name):
    tail = clean(tail)
    variant_name = clean(variant_name)
    if not tail:
        return variant_name
    if variant_name.upper().startswith(tail.upper()):
        return variant_name
    return f"{tail} / {variant_name}" if variant_name else tail


def numeric_range(values, multiplier=None):
    numbers = [float(value) for value in values]
    if multiplier is None:
        return min(numbers), max(numbers)
    priced = [math.ceil(number * multiplier) for number in numbers]
    return min(priced), max(priced)


def variant_sku_for(product_preview_id, variant, index):
    sku_part = slugify(variant.get("source_code") or variant.get("variant_name") or f"variant-{index}").upper()
    return f"OK-{sku_part}-{stable_id(product_preview_id, variant.get('variant_name'), index).upper()}"[:96]


def split_large_products_by_motorcycle_brand(products, variants_by_product):
    split_products = []
    split_variants = []
    split_product_count = 0
    created_brand_products = 0

    for product in products:
        product_variants = variants_by_product.get(product["preview_id"], [])
        if len(product_variants) < BRAND_SPLIT_VARIANT_THRESHOLD:
            split_products.append(product)
            split_variants.extend(product_variants)
            continue

        grouped = {}
        for variant in product_variants:
            grouped.setdefault(motorcycle_brand_for_variant(variant), []).append(variant)

        recognized_brands = [brand for brand in grouped if brand != "Other Fitments"]
        recognized_variant_count = sum(len(grouped[brand]) for brand in recognized_brands)
        if len(recognized_brands) < 2 or recognized_variant_count < BRAND_SPLIT_MIN_RECOGNIZED_VARIANTS:
            split_products.append(product)
            split_variants.extend(product_variants)
            continue

        split_product_count += 1
        created_brand_products += len(grouped)
        base_name = clean(product["english_product_name"])
        base_title = clean(product["source_title"])
        for brand in sorted(grouped, key=lambda value: (value == "Other Fitments", value)):
            brand_variants = grouped[brand]
            brand_preview_id = f"{product['preview_id']}-{slugify(brand)}"
            costs = [float(variant["cost"]) for variant in brand_variants]
            cost_min, cost_max = numeric_range(costs)
            retail_min, retail_max = numeric_range(costs, 1.20)
            wholesale_min, wholesale_max = numeric_range(costs, 1.12)
            image_url = next(
                (variant.get("source_image_url") for variant in brand_variants if clean(variant.get("source_image_url"))),
                clean(product.get("main_image_url")),
            )
            note_prefix = clean(product.get("review_note"))
            split_note = f"Split from {base_name} by motorcycle brand. Vehicle models stay as variants."
            if note_prefix:
                split_note = f"{note_prefix}; {split_note}"

            draft_product = dict(product)
            draft_product.update(
                {
                    "preview_id": brand_preview_id,
                    "source_title": f"{base_title} - {brand}",
                    "english_product_name": f"{base_name} - {brand}",
                    "variant_count": len(brand_variants),
                    "cost_min": cost_min,
                    "cost_max": cost_max,
                    "retail_min_1_5pcs": retail_min,
                    "retail_max_1_5pcs": retail_max,
                    "wholesale_min_6pcs_plus": wholesale_min,
                    "wholesale_max_6pcs_plus": wholesale_max,
                    "main_image_url": image_url,
                    "needs_review": "NO" if image_url else "YES",
                    "review_note": split_note + ("" if image_url else "; missing image"),
                }
            )
            split_products.append(draft_product)

            for index, variant in enumerate(brand_variants, start=1):
                draft_variant = dict(variant)
                draft_variant["preview_id"] = brand_preview_id
                draft_variant["variant_sku"] = variant_sku_for(brand_preview_id, draft_variant, index)
                draft_variant["sort_order"] = index
                split_variants.append(draft_variant)

    return split_products, split_variants, split_product_count, created_brand_products


def write_xlsx(path, products, variants):
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = "Products Merged"
    product_fields = list(products[0].keys()) if products else []
    ws.append(product_fields)
    for row in products:
        ws.append([row.get(field, "") for field in product_fields])
    ws2 = workbook.create_sheet("Variants Merged")
    variant_fields = list(variants[0].keys()) if variants else []
    ws2.append(variant_fields)
    for row in variants:
        ws2.append([row.get(field, "") for field in variant_fields])
    ws3 = workbook.create_sheet("Missing Images")
    ws3.append(product_fields)
    for row in products:
        if row.get("needs_review") == "YES":
            ws3.append([row.get(field, "") for field in product_fields])
    workbook.save(path)


def main():
    products = read_csv(PRODUCTS_CSV)
    variants = read_csv(VARIANTS_CSV)
    variants_by_product = {}
    for variant in variants:
        variants_by_product.setdefault(variant["preview_id"], []).append(variant)

    candidates = {}
    product_meta = {}
    for product in products:
        base, tail = model_base(product["source_title"], product["sheet"])
        key = (product["sheet"], base)
        product_meta[product["preview_id"]] = {"base": base, "tail": tail, "key": key}
        candidates.setdefault(key, []).append(product)

    new_products = []
    new_variants = []
    merged_groups = 0
    merged_original_products = 0
    old_to_new = {}

    for key, group in candidates.items():
        sheet, base = key
        merge = should_merge(group, product_meta)
        if merge:
            merged_groups += 1
            merged_original_products += len(group)
            preview_id = f"ok-models-{slugify(sheet)}-{stable_id(sheet, base)}"
            all_variants = []
            for product in group:
                meta = product_meta[product["preview_id"]]
                old_to_new[product["preview_id"]] = preview_id
                for variant in variants_by_product.get(product["preview_id"], []):
                    draft = dict(variant)
                    draft["preview_id"] = preview_id
                    draft["variant_name"] = append_tail_to_variant(meta["tail"], draft["variant_name"])
                    draft["model"] = clean(" / ".join(part for part in [meta["tail"], draft.get("model", "")] if clean(part)))
                    all_variants.append(draft)
            costs = [float(variant["cost"]) for variant in all_variants]
            retail = [math.ceil(cost * 1.20) for cost in costs]
            wholesale = [math.ceil(cost * 1.12) for cost in costs]
            image_url = next((product["main_image_url"] for product in group if clean(product["main_image_url"])), "")
            source_titles = "; ".join(product["source_title"] for product in group[:12])
            if len(group) > 12:
                source_titles += f"; +{len(group) - 12} more"
            template = dict(group[0])
            template.update(
                {
                    "preview_id": preview_id,
                    "source_title": base,
                    "english_product_name": merged_product_name(base),
                    "variant_count": len(all_variants),
                    "cost_min": min(costs),
                    "cost_max": max(costs),
                    "retail_min_1_5pcs": min(retail),
                    "retail_max_1_5pcs": max(retail),
                    "wholesale_min_6pcs_plus": min(wholesale),
                    "wholesale_max_6pcs_plus": max(wholesale),
                    "main_image_url": image_url,
                    "needs_review": "NO" if image_url else "YES",
                    "review_note": (
                        f"Merged {len(group)} model/fitment products into variants. Original titles: {source_titles}"
                        + ("" if image_url else "; missing image")
                    ),
                }
            )
            new_products.append(template)
            new_variants.extend(all_variants)
        else:
            for product in group:
                new_products.append(product)
                new_variants.extend(variants_by_product.get(product["preview_id"], []))

    final_variants_by_product = {}
    for variant in new_variants:
        final_variants_by_product.setdefault(variant["preview_id"], []).append(variant)
    new_products, new_variants, brand_split_products, brand_products_created = split_large_products_by_motorcycle_brand(
        new_products,
        final_variants_by_product,
    )

    product_fields = list(products[0].keys())
    variant_fields = list(variants[0].keys())
    write_csv(PREVIEW_DIR / "ok-products-preview-merged-models.csv", new_products, product_fields)
    write_csv(PREVIEW_DIR / "ok-variants-preview-merged-models.csv", new_variants, variant_fields)
    workbook_path = PREVIEW_DIR / "ok-upload-preview-merged-models.xlsx"
    try:
        write_xlsx(workbook_path, new_products, new_variants)
    except PermissionError:
        workbook_path = PREVIEW_DIR / f"ok-upload-preview-merged-models-{stable_id(len(new_products), len(new_variants))}.xlsx"
        write_xlsx(workbook_path, new_products, new_variants)

    summary = {
        "original_products": len(products),
        "original_variants": len(variants),
        "merged_products": len(new_products),
        "merged_variants": len(new_variants),
        "merged_groups": merged_groups,
        "merged_original_products": merged_original_products,
        "brand_split_source_products": brand_split_products,
        "brand_split_created_products": brand_products_created,
        "brand_split_threshold": BRAND_SPLIT_VARIANT_THRESHOLD,
        "products_with_images": sum(1 for product in new_products if clean(product.get("main_image_url"))),
        "products_missing_images": sum(1 for product in new_products if not clean(product.get("main_image_url"))),
        "pricing_rule": "cost=PRICE from supplier sheet, retail=ceil(cost*1.20), 6pcs_plus=ceil(cost*1.12)",
        "files": {
            "products_csv": str(PREVIEW_DIR / "ok-products-preview-merged-models.csv"),
            "variants_csv": str(PREVIEW_DIR / "ok-variants-preview-merged-models.csv"),
            "workbook": str(workbook_path),
            "summary_json": str(PREVIEW_DIR / "ok-merged-models-summary.json"),
        },
    }
    (PREVIEW_DIR / "ok-merged-models-summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
