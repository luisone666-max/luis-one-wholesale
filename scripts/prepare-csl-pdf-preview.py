import csv
import hashlib
import json
import math
import re
from collections import defaultdict
from pathlib import Path

import fitz
import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = Path("C:/Users/Administrator/Desktop/\u4e8c\u7ef4\u7801/csl")
OUT_DIR = ROOT / "tmp-imports" / "csl-preview"
PUBLIC_IMAGE_DIR = ROOT / "public" / "product-images" / "csl-preview"

FILES = [
    ("CSL", SOURCE_DIR / "CSL \u4ef7\u683c\u8868 \u5185\u914d 2-12-2026(1).pdf"),
    ("CSL", SOURCE_DIR / "CSL \u8f6e\u80ce\uff0c\u5185\u80ce\uff0c\u771f\u7a7a\u80ce 4-7-2026(1).pdf"),
    ("CSL", SOURCE_DIR / "CSL \u88c5\u9970\u54c1\uff0c\u5916\u88c5\u4ea7\u54c1  4-7-2026.pdf"),
    ("CSL", SOURCE_DIR / "\u62a4\u7406\u4ea7\u54c1\uff0c\u6392\u6c14\u7ba1 4-8-2026(1).pdf"),
    ("CRISTAL", SOURCE_DIR / "CRISTAL \u94fe\u997c \u7ebf\u6761 \u4ef7\u683c\u8868 1-6-2026.pdf"),
    ("CSL", SOURCE_DIR / "CSL \u4ef7\u683c\u8868 \u9a6c\u5c3c\u62c9 4-20-2026.pdf"),
]

BAD_TITLE_WORDS = {
    "COVER",
    "CATALOG",
    "PAGE",
    "PRODUCT",
}

EXCLUDE_TERMS = {
    "STICKER",
    "DECAL",
    "TANK PROTECTION",
    "ADHESIVE",
    "ALUMINUM NAME",
    "REFLECTIVE STICKER",
}

GENERIC_BASES = {
    "ACCESSORIES",
    "ACCESSORY",
    "PARTS",
    "MAINTENANCE PRODUCT",
    "PLASTIC ACCESSORIES",
}

COLORS = {
    "BLACK",
    "BLCK",
    "BLUE",
    "RED",
    "GOLD",
    "SILVER",
    "ORANGE",
    "WHITE",
    "GREEN",
    "YELLOW",
    "PINK",
    "GRAY",
    "GREY",
    "CHROME",
    "TITANIUM",
    "STAINLESS",
    "CARBON",
    "BROWN",
    "PURPLE",
    "CLEAR",
    "NEON",
}

COLOR_ALIASES = {
    "BLCK": "BLACK",
    "SIL": "SILVER",
}

QTY_TOKEN_RE = re.compile(
    r"(?i)(?:\b\d+\s*(?:PCS|PC|SET|SETS|PAIR|PAIRS|ROLL|ROLLS|BOTTLE|BOTTLES|CAN|CANS|BOX(?:ES)?|PACKS?)\b|/[ ]?(?:BOX(?:ES)?|PACKS?)\b|\b(?:BOX(?:ES)?|PACKS?)\b)"
)

MODEL_TOKEN_BRANDS = {
    "HONDA": "Honda",
    "YAMAHA": "Yamaha",
    "SUZUKI": "Suzuki",
    "KAWASAKI": "Kawasaki/Bajaj",
    "BAJAJ": "Kawasaki/Bajaj",
    "ADV": "Honda",
    "ADV150": "Honda",
    "ADV160": "Honda",
    "BEAT": "Honda",
    "BEAT110": "Honda",
    "BEATFI": "Honda",
    "C100": "Honda",
    "CB110": "Honda",
    "CB125": "Honda",
    "CG125": "Honda",
    "CG150": "Honda",
    "CLICK": "Honda",
    "CLICK125": "Honda",
    "CLICK150": "Honda",
    "CLICK160": "Honda",
    "DASH": "Honda",
    "DASH110": "Honda",
    "DIO": "Honda",
    "DREAM": "Honda",
    "PCX": "Honda",
    "PCX150": "Honda",
    "PCX160": "Honda",
    "RS100": "Honda",
    "RS110": "Honda",
    "RS125": "Honda",
    "RS150": "Honda",
    "SCOOPY": "Honda",
    "SCOOPYFI": "Honda",
    "SONIC": "Honda",
    "SONIC150": "Honda",
    "STX": "Honda",
    "STX125": "Honda",
    "SUPREMO": "Honda",
    "TIGER": "Honda",
    "TMX": "Honda",
    "TMX125": "Honda",
    "TMX155": "Honda",
    "WAVE": "Honda",
    "WAVE100": "Honda",
    "WAVE100R": "Honda",
    "WAVE110": "Honda",
    "WAVE125": "Honda",
    "WAVE125I": "Honda",
    "WAVE125R": "Honda",
    "XLR200": "Honda",
    "XR200": "Honda",
    "XRM": "Honda",
    "XRM110": "Honda",
    "XRM125": "Honda",
    "AEROX": "Yamaha",
    "AEROX155": "Yamaha",
    "CRYPTON": "Yamaha",
    "F16": "Yamaha",
    "FZ16": "Yamaha",
    "GRAVIS": "Yamaha",
    "LC135": "Yamaha",
    "LC150": "Yamaha",
    "MIO": "Yamaha",
    "MIO110": "Yamaha",
    "MIO125": "Yamaha",
    "NMAX": "Yamaha",
    "NMAX155": "Yamaha",
    "NOUVO": "Yamaha",
    "SZ16": "Yamaha",
    "SNIPER": "Yamaha",
    "SNIPER135": "Yamaha",
    "SNIPER150": "Yamaha",
    "SNIPER155": "Yamaha",
    "SOUL": "Yamaha",
    "VEGA": "Yamaha",
    "X1": "Yamaha",
    "X1R": "Yamaha",
    "XTZ125": "Yamaha",
    "YTX": "Yamaha",
    "YTX125": "Yamaha",
    "AXELO": "Suzuki",
    "BURGMAN": "Suzuki",
    "GD110": "Suzuki",
    "GP125": "Suzuki",
    "GS125": "Suzuki",
    "NEX": "Suzuki",
    "RAIDER": "Suzuki",
    "RAIDER150": "Suzuki",
    "RAIDERJ110": "Suzuki",
    "SHOGUN": "Suzuki",
    "SHOGUN110": "Suzuki",
    "SKYDRIVE": "Suzuki",
    "SMASH": "Suzuki",
    "SMASH110": "Suzuki",
    "SMASH115": "Suzuki",
    "BARAKO": "Kawasaki/Bajaj",
    "BARAKO175": "Kawasaki/Bajaj",
    "BC175": "Kawasaki/Bajaj",
    "CT100": "Kawasaki/Bajaj",
    "CT125": "Kawasaki/Bajaj",
    "CT150": "Kawasaki/Bajaj",
    "FURY": "Kawasaki/Bajaj",
    "FURY125": "Kawasaki/Bajaj",
    "HD3": "Kawasaki/Bajaj",
    "HDIII": "Kawasaki/Bajaj",
    "ROUSER": "Kawasaki/Bajaj",
    "WIND125": "Kawasaki/Bajaj",
    "RUSI": "Rusi",
    "RUSI125": "Rusi",
    "RUSI150": "Rusi",
    "RUSI175": "Rusi",
    "GY6": "GY6 Scooter",
    "LIFAN": "Lifan",
    "LIFAN110": "Lifan",
    "LIFAN125": "Lifan",
    "LIFAN150": "Lifan",
}


def clean(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\t", " ").strip())


def ascii_clean(value):
    text = clean(value)
    text = re.sub(r"[^\x00-\x7f]+", " ", text)
    return clean(text)


def slugify(value, fallback="item"):
    text = ascii_clean(value).lower()
    text = re.sub(r"[^a-z0-9]+", "-", text)
    return re.sub(r"-+", "-", text).strip("-") or fallback


def stable_id(*parts):
    raw = "|".join(clean(part) for part in parts)
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:10]


def title_case_codes(value):
    words = []
    for part in ascii_clean(value).split():
        if re.search(r"\d", part) or part.isupper() or "/" in part or "-" in part:
            words.append(part.upper())
        else:
            words.append(part.capitalize())
    return " ".join(words)


def price_value(line):
    text = clean(line).replace(",", "")
    match = re.search(r"₱\s*(\d+(?:\.\d+)?)", text)
    if not match:
        return None
    return float(match.group(1))


def is_price_line(line):
    return price_value(line) is not None


def is_footer(line):
    return bool(re.match(r"^第\s*\d+\s*页", clean(line)))


def has_excluded_term(text):
    upper = clean(text).upper()
    return any(term in upper for term in EXCLUDE_TERMS)


def normalize_title(raw_title):
    title = ascii_clean(raw_title).upper()
    title = title.replace("SROCKET", "SPROCKET")
    title = re.sub(r"\*.*?NET PRICE.*?\*", " ", title)
    title = title.replace("NET PRICE", " ")
    title = title.replace("GENUINE PARTS", " ")
    title = re.sub(r"\bCSL\s+NEW\s+TOPLUCK\s+TRADING\b", " ", title)
    title = re.sub(r"#\s*\d+[A-Z]?", " ", title)
    title = re.sub(r"\s+#\s*", " ", title)
    title = title.replace("#", " ")
    title = re.sub(r"\bNO\.\s*\d+\b", " ", title)
    title = re.sub(r"\b[1-9]\b$", " ", title)
    title = re.sub(r"\b\d+\b$", " ", title)
    title = re.sub(r"\([^A-Z0-9]*\)", " ", title)
    title = re.sub(r"[^A-Z0-9/&()# .-]+", " ", title)
    title = clean(title)
    replacements = {
        "CSL TYRE 6LAYER": "6-Layer Tube Type Tire",
        "CSL TUBELESS TYRE 6LAYER": "6-Layer Tubeless Tire",
        "CSL INTERIOR TUBE": "Inner Tube",
        "APIDO PIPE": "APIDO Exhaust Pipe",
        "APIDO CHICKEN PIPE": "APIDO Chicken Pipe Exhaust",
        "BOSNY SPRAY PAINT 400ML": "BOSNY Spray Paint 400ML",
        "SAMURAI SPRAY PAINT 400ML": "SAMURAI Spray Paint 400ML",
        "HANDLE LEVER CNC": "CNC Handle Lever",
        "LEVER TWO TONE": "Two Tone Lever",
        "FANCOVER": "Fan Cover",
        "HEADLAMP": "Headlamp",
        "WINKERLAMP": "Winker Lamp",
        "TAIL LAMP": "Tail Lamp",
        "BOLT NUT": "Bolt Nut",
        "BOLT-NUT": "Bolt Nut",
        "STOCK LEVER": "Stock Lever",
        "ROTOR DISC": "Rotor Disc",
        "CHANGE GEAR PEDAL": "Change Gear Pedal",
        "KICKSTARTER ARM ASSY": "Kickstarter Arm Assy",
        "MONORACK": "Monorack Bracket",
    }
    for source, target in replacements.items():
        if title.startswith(source):
            suffix = clean(title[len(source) :])
            suffix = clean(re.sub(r"^\d+\b|\b\d+$", " ", suffix))
            return clean(f"{target} {suffix}")
    return title_case_codes(title)


def detect_supplier(default_supplier, title):
    upper = clean(title).upper()
    for supplier in ["APIDO", "BOSNY", "SAMURAI", "CRISTAL", "CSL"]:
        if upper.startswith(supplier) or f" {supplier} " in f" {upper} ":
            return supplier
    return default_supplier


def is_qty(line):
    text = clean(line).upper().replace(" ", "")
    return bool(QTY_TOKEN_RE.fullmatch(text))


def split_qty_from_line(line):
    text = clean(line)
    qty_parts = [clean(match.group(0).upper().replace(" ", "")).lstrip("/") for match in QTY_TOKEN_RE.finditer(text)]
    remainder = clean(QTY_TOKEN_RE.sub(" ", text))
    return remainder, qty_parts


def merge_qty_parts(*values):
    parts = []
    for value in values:
        for part in clean(value).split(" / "):
            part = clean(part)
            if part and part not in parts:
                parts.append(part)
    return " / ".join(parts)


def normalize_color_token(token):
    token = clean(token).upper().replace(" ", "")
    token = COLOR_ALIASES.get(token, token)
    return token if token in COLORS else ""


def colors_from_line(line):
    text = clean(line).upper()
    text = text.replace("COLOCR", "").replace("COLOR", "")
    tokens = re.split(r"[/,; ]+", text)
    colors = [normalize_color_token(token) for token in tokens]
    return [color for color in colors if color]


def is_code(line):
    text = clean(line).upper()
    if text.startswith("#"):
        return True
    if re.search(r"\b(?:CSL|LT|OK|RP|B|G|S|T)-[A-Z0-9-]+\b", text):
        return True
    return False


def is_size_or_spec(line):
    text = clean(line).upper()
    if text in {"STD", "STANDARD", "OFF", "ROAD"}:
        return True
    if re.fullmatch(r"\d+(?:\.\d+)?MM", text):
        return True
    if re.fullmatch(r"\d{2,3}T", text):
        return True
    if re.fullmatch(r"\d+(?:/\d+)?-\d+[A-Z]*", text):
        return True
    if re.fullmatch(r"\d+[A-Z]-\d+[A-Z]*", text):
        return True
    if re.fullmatch(r"\d+(?:\.\d+)?OZ", text):
        return True
    if "LAYER" in text:
        return True
    return False


def tokens_for_brand_detection(value):
    tokens = []
    text = clean(value).upper()
    text = re.sub(r"N[\s-]?MAX", "NMAX", text)
    text = re.sub(r"SKY\s*DRIVE", "SKYDRIVE", text)
    text = re.sub(r"BEAT\s*FI", "BEATFI", text)
    text = re.sub(r"SCOOPY\s*FI", "SCOOPYFI", text)
    text = re.sub(r"RAIDER\s*J\s*110", "RAIDERJ110", text)
    text = text.replace("AREOX", "AEROX")
    for token in re.split(r"[^A-Z0-9]+", text):
        token = re.sub(r"[^A-Z0-9]", "", token)
        if token:
            tokens.append(token)
    return tokens


def motorcycle_brand_for_variant(variant):
    values = [
        variant.get("model", ""),
        variant.get("variant_name", ""),
        variant.get("size", ""),
        variant.get("source_code", ""),
    ]
    brands = []
    for token in tokens_for_brand_detection(" ".join(values)):
        brand = MODEL_TOKEN_BRANDS.get(token)
        if brand and brand not in brands:
            brands.append(brand)
    if not brands:
        return "Other Fitments"
    if "Kawasaki/Bajaj" in brands:
        return "Kawasaki/Bajaj"
    return brands[0]


def category_for_title(title):
    upper = clean(title).upper()
    if "BULB" in upper:
        return "Motorcycle Parts > Lights & Electrical > BULB"
    if any(term in upper for term in ["SOCKET", "SWITCH", "CDI", "STATOR", "REGULATOR", "COIL", "RELAY", "SENSOR", "STARTER", "MAGNETO"]):
        return "Motorcycle Parts > Lights & Electrical > Electrical Parts"
    if any(term in upper for term in ["SPARK PLUG", "SPARKPLUG", "SPARPLUG"]):
        return "Motorcycle Parts > Engine Parts > SPARK PLUG"
    if any(term in upper for term in ["FILTER", "AIR CLEANER"]):
        return "Motorcycle Parts > Engine Parts > FILTER"
    if any(term in upper for term in ["HUB", "AXLE", "BEARING", "SPOKE", "RIM", "WHEEL"]):
        return "Motorcycle Parts > Wheels & Tires > Wheels & Spokes"
    if any(term in upper for term in ["BOLT", "NUT", "SCREW"]):
        return "Motorcycle Parts > Motorcycle Accessories > Hardware"
    if any(term in upper for term in ["TIRE", "TYRE", "TUBE", "RIM", "SPOKE"]):
        if "PITO" in upper or "VALVE" in upper:
            return "Motorcycle Parts > Wheels & Tires > Tire Accessories"
        if "TUBELESS" in upper or "TIRE" in upper or "TYRE" in upper:
            return "Motorcycle Parts > Wheels & Tires > Tires"
        if "TUBE" in upper:
            return "Motorcycle Parts > Wheels & Tires > Inner Tubes"
        return "Motorcycle Parts > Wheels & Tires > Wheels & Spokes"
    if any(term in upper for term in ["CHAIN", "SPROCKET"]):
        return "Motorcycle Parts > Drive Train > Chain & Sprocket"
    if "CABLE" in upper:
        return "Motorcycle Parts > Controls & Cables > Cables"
    if any(term in upper for term in ["BRAKE", "CALIPER", "ROTOR DISC", "DISC"]):
        return "Motorcycle Parts > Brake System > Brake Parts"
    if any(term in upper for term in ["HEADLAMP", "HEADLIGHT", "LED", "WINKER", "TAIL LAMP", "LAMP", "HORN"]):
        return "Motorcycle Parts > Lights & Electrical > Lights & Horns"
    if any(term in upper for term in ["PISTON", "CYLINDER", "VALVE", "GASKET", "BEARING", "OIL SEAL", "O-RING", "ROCKER", "CAM", "CARBURETOR", "MANIFOLD", "CLUTCH"]):
        return "Motorcycle Parts > Engine Parts > Engine Components"
    if any(term in upper for term in ["PIPE", "EXHAUST", "MUFFLER"]):
        return "Motorcycle Parts > Exhaust > Exhaust Parts"
    if any(term in upper for term in ["SPRAY", "WD-40", "BRAKE FLUID", "CLEANER", "RUST", "GREASE", "OIL"]):
        return "Motorcycle Parts > Care & Maintenance > Maintenance Products"
    if any(term in upper for term in ["SEAT", "FENDER", "COVER", "FANCOVER", "TANK", "BODY"]):
        return "Motorcycle Parts > Body Parts > Body Covers"
    if any(term in upper for term in ["LEVER", "HANDLE", "FOOTREST", "PEDAL", "STAND", "MIRROR", "GRIP"]):
        return "Motorcycle Parts > Motorcycle Accessories > Controls & Accessories"
    if any(term in upper for term in ["BRACKET", "MONORACK", "GUARD", "HOOK", "AXLE CAP"]):
        return "Motorcycle Parts > Motorcycle Accessories > Brackets & Guards"
    return "Motorcycle Parts > Motorcycle Accessories > General Accessories"


def public_path_for_image(file_stem, page_index, image_index, ext):
    return f"/product-images/csl-preview/{slugify(file_stem)}/p{page_index + 1:03d}-img{image_index + 1:02d}{ext}"


def extract_page_images(doc, page, file_stem, page_index):
    images = []
    for image_index, info in enumerate(page.get_image_info(xrefs=True)):
        xref = info.get("xref")
        if not xref:
            continue
        try:
            extracted = doc.extract_image(xref)
            data = extracted["image"]
            ext = "." + (extracted.get("ext") or "png").lower()
        except Exception:
            pix = page.get_pixmap(clip=fitz.Rect(info["bbox"]), matrix=fitz.Matrix(2, 2), alpha=False)
            data = pix.tobytes("png")
            ext = ".png"
        out_public = public_path_for_image(file_stem, page_index, image_index, ext)
        out_path = ROOT / "public" / out_public.lstrip("/")
        out_path.parent.mkdir(parents=True, exist_ok=True)
        if not out_path.exists() or out_path.read_bytes() != data:
            out_path.write_bytes(data)
        bbox = info["bbox"]
        images.append(
            {
                "path": out_public,
                "x0": bbox[0],
                "y0": bbox[1],
                "x1": bbox[2],
                "y1": bbox[3],
                "cy": (bbox[1] + bbox[3]) / 2,
            }
        )
    return images


def price_y_positions(page):
    positions = []
    for word in sorted(page.get_text("words"), key=lambda item: (item[1], item[0])):
        if is_price_line(word[4]):
            positions.append((word[4], (word[1] + word[3]) / 2))
    return positions


def image_for_y(images, y):
    if not images:
        return ""
    left_images = [image for image in images if image["x0"] < 220]
    candidates = left_images or images
    containing = [image for image in candidates if image["y0"] - 12 <= y <= image["y1"] + 12]
    if containing:
        return min(containing, key=lambda image: abs(image["cy"] - y))["path"]
    return min(candidates, key=lambda image: abs(image["cy"] - y))["path"]


def title_and_data_lines(lines):
    header_index = None
    for index, line in enumerate(lines):
        if clean(line).upper() == "IMAGE":
            header_index = index
            break
    if header_index is None:
        return "", []
    title_lines = []
    for line in lines[:header_index]:
        upper = clean(line).upper()
        if upper in {"CSL NEW TOPLUCK TRADING", "CRISTAL", "GENUINE PARTS"}:
            continue
        if "ADDRESS:" in upper or "CONTACT:" in upper:
            continue
        if upper.startswith("(") and "NET" in upper:
            continue
        title_lines.append(line)
    raw_title = clean(" ".join(title_lines))
    header_end = header_index
    for index in range(header_index + 1, min(len(lines), header_index + 10)):
        upper = clean(lines[index]).upper()
        if upper in {"MODELS", "MODEL", "COLOR", "COLOCR", "SIZE", "QTY/BOX", "PCS/BOX", "PRICE", "PRICE/SET", "COLOR/PCS", "MOTORCYCLE MODEL", "CODE", "OE"}:
            header_end = index
            continue
        break
    return raw_title, lines[header_end + 1 :]


def parse_segment(segment):
    qty = ""
    colors = []
    codes = []
    sizes = []
    remaining = []
    for raw in segment:
        line = clean(raw)
        if not line or is_footer(line):
            continue
        line, inline_qty = split_qty_from_line(line)
        if inline_qty:
            qty = merge_qty_parts(qty, " / ".join(inline_qty))
        if not line:
            continue
        if is_qty(line):
            qty = merge_qty_parts(qty, line.upper().replace(" ", ""))
            continue
        found_colors = colors_from_line(line)
        if found_colors and len(found_colors) == len(re.split(r"[/,; ]+", line.strip())):
            colors.extend(found_colors)
            continue
        if is_code(line):
            codes.append(line)
            continue
        if is_size_or_spec(line):
            sizes.append(line)
            continue
        remaining.append(line)

    if not colors:
        colors = [""]
    colors = list(dict.fromkeys(COLOR_ALIASES.get(color, color) for color in colors))
    source_code = " / ".join(codes)
    size = " ".join(sizes)
    model = clean(" ".join(remaining))
    if not model and size:
        model, size = size, ""
    return model, size, colors, qty, source_code


def page_rows(default_supplier, path, doc, page_index, excluded):
    page = doc[page_index]
    lines = [clean(line) for line in page.get_text("text").splitlines() if clean(line)]
    raw_title, data_lines = title_and_data_lines(lines)
    if not raw_title:
        return []
    title = normalize_title(raw_title)
    supplier = detect_supplier(default_supplier, raw_title)
    if has_excluded_term(raw_title):
        excluded.append(
            {
                "source_file": path.name,
                "source_page": page_index + 1,
                "reason": "Excluded sticker/decal page",
                "raw_title": raw_title,
            }
        )
        return []

    images = extract_page_images(doc, page, path.stem, page_index)
    price_positions = price_y_positions(page)
    price_index = 0
    rows = []
    segment = []
    last_qty = ""
    for line in data_lines:
        if is_footer(line):
            break
        if has_excluded_term(line):
            excluded.append(
                {
                    "source_file": path.name,
                    "source_page": page_index + 1,
                    "reason": "Excluded sticker/decal row",
                    "raw_title": raw_title,
                    "row_text": line,
                }
            )
            segment = []
            continue
        price = price_value(line)
        if price is None:
            segment.append(line)
            continue
        model, size, colors, qty, source_code = parse_segment(segment)
        if not qty:
            qty = last_qty
        if qty:
            last_qty = qty
        price_y = price_positions[price_index][1] if price_index < len(price_positions) else 400
        price_index += 1
        image_url = image_for_y(images, price_y)
        for color in colors:
            rows.append(
                {
                    "supplier": supplier,
                    "source_file": path.name,
                    "source_page": page_index + 1,
                    "raw_title": raw_title,
                    "product_base": title,
                    "model": model,
                    "size": size,
                    "color": color,
                    "qty_box": qty,
                    "source_code": source_code,
                    "cost": price,
                    "source_image_url": image_url,
                    "price_y": price_y,
                }
            )
        segment = []
    return rows


def should_use_row_as_base(base):
    upper = clean(base).upper()
    if upper in GENERIC_BASES:
        return True
    if re.fullmatch(r"PARTS(?: \d+)?", upper):
        return True
    if re.fullmatch(r"ACCESSORIES(?: \d+)?", upper):
        return True
    if upper.startswith("MAINTENANCE PRODUCT"):
        return True
    if upper.startswith("PLASTIC ACCESSORIES"):
        return True
    return False


def row_base(row):
    base = clean(row["product_base"])
    if should_use_row_as_base(base) and clean(row.get("model")):
        model = clean(row["model"])
        model = re.sub(r"\s*/\s*(?:PACK|PACKS|BOX|BOXES)\b", " ", model, flags=re.I)
        model = clean(QTY_TOKEN_RE.sub(" ", model))
        if len(model.split()) <= 6:
            return normalize_title(model)
    return base


def variant_label(row):
    parts = [row.get("model", ""), row.get("size", ""), row.get("color", ""), row.get("source_code", "")]
    return clean(" / ".join(part for part in parts if clean(part))) or "Default"


def product_display_name(supplier, base, brand_suffix=""):
    base = clean(base)
    name_base = title_case_codes(base)
    if name_base.upper().startswith(supplier.upper() + " "):
        name = name_base
    else:
        name = f"{supplier} {name_base}"
    if brand_suffix:
        name = f"{name} - {brand_suffix}"
    return clean(name)


def split_group_key(supplier, base, variants):
    brand_counts = defaultdict(int)
    for variant in variants:
        brand_counts[motorcycle_brand_for_variant(variant)] += 1
    recognized = [brand for brand in brand_counts if brand != "Other Fitments"]
    recognized_count = sum(brand_counts[brand] for brand in recognized)
    if len(variants) >= 8 and len(recognized) >= 2 and recognized_count >= 4:
        groups = defaultdict(list)
        for variant in variants:
            groups[motorcycle_brand_for_variant(variant)].append(variant)
        return groups
    return {"": variants}


def dedupe_variants(rows):
    seen = {}
    for row in rows:
        key = (
            row["supplier"],
            row_base(row).upper(),
            clean(row.get("model")).upper(),
            clean(row.get("size")).upper(),
            clean(row.get("color")).upper(),
            clean(row.get("source_code")).upper(),
            round(float(row["cost"]), 2),
        )
        # Later dated PDFs should win when exact duplicates exist, so overwrite.
        seen[key] = row
    return list(seen.values())


def build_preview(rows):
    rows = dedupe_variants(rows)
    grouped = defaultdict(list)
    for row in rows:
        base = row_base(row)
        grouped[(row["supplier"], base)].append(row)

    products = []
    variants = []
    for (supplier, base), group_rows in sorted(grouped.items(), key=lambda item: (item[0][0], item[0][1])):
        split_groups = split_group_key(supplier, base, group_rows)
        for brand_suffix, variant_rows in sorted(split_groups.items(), key=lambda item: (item[0] == "Other Fitments", item[0])):
            preview_id = f"csl-{slugify(supplier)}-{slugify(base)}-{slugify(brand_suffix or 'all')}-{stable_id(supplier, base, brand_suffix)}"
            costs = [float(row["cost"]) for row in variant_rows]
            retail = [math.ceil(cost * 1.20) for cost in costs]
            wholesale = [math.ceil(cost * 1.12) for cost in costs]
            image_url = next((row["source_image_url"] for row in variant_rows if clean(row.get("source_image_url"))), "")
            source_pages = sorted({f"{row['source_file']} p.{row['source_page']}" for row in variant_rows})
            source_titles = sorted({clean(row["raw_title"]) for row in variant_rows})
            product = {
                "preview_id": preview_id,
                "supplier": supplier,
                "source_files_pages": "; ".join(source_pages[:20]) + (f"; +{len(source_pages) - 20} more" if len(source_pages) > 20 else ""),
                "source_title": "; ".join(source_titles[:8]) + (f"; +{len(source_titles) - 8} more" if len(source_titles) > 8 else ""),
                "category": category_for_title(base),
                "english_product_name": product_display_name(supplier, base, brand_suffix),
                "variant_count": len(variant_rows),
                "cost_min": min(costs),
                "cost_max": max(costs),
                "retail_min_1_5pcs": min(retail),
                "retail_max_1_5pcs": max(retail),
                "wholesale_min_6pcs_plus": min(wholesale),
                "wholesale_max_6pcs_plus": max(wholesale),
                "main_image_url": image_url,
                "needs_review": "YES" if not image_url or any(not clean(row.get("model")) for row in variant_rows) else "NO",
                "review_note": (
                    ("missing image; " if not image_url else "")
                    + ("some rows missing model/spec; " if any(not clean(row.get("model")) for row in variant_rows) else "")
                    + ("split by motorcycle brand; " if brand_suffix else "")
                    + "PDF price treated as supplier net cost."
                ).strip(),
            }
            products.append(product)

            for index, row in enumerate(variant_rows, start=1):
                label = variant_label(row)
                variants.append(
                    {
                        "preview_id": preview_id,
                        "variant_sku": f"CSL-{slugify(supplier).upper()}-{stable_id(preview_id, label, index).upper()}",
                        "variant_name": label,
                        "model": clean(row.get("model")),
                        "color": clean(row.get("color")),
                        "size": clean(row.get("size")),
                        "source_code": clean(row.get("source_code")),
                        "qty_box": clean(row.get("qty_box")),
                        "cost": float(row["cost"]),
                        "retail_price_1_5pcs": math.ceil(float(row["cost"]) * 1.20),
                        "wholesale_price_6pcs_plus": math.ceil(float(row["cost"]) * 1.12),
                        "source_image_url": clean(row.get("source_image_url")),
                        "source_file": row["source_file"],
                        "source_page": row["source_page"],
                        "source_title": row["raw_title"],
                        "sort_order": index,
                    }
                )
    return products, variants


def write_csv(path, rows, fieldnames):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(path, products, variants, excluded):
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = "Products Preview"
    product_fields = list(products[0].keys()) if products else []
    ws.append(product_fields)
    for row in products:
        ws.append([row.get(field, "") for field in product_fields])

    ws2 = workbook.create_sheet("Variants Preview")
    variant_fields = list(variants[0].keys()) if variants else []
    ws2.append(variant_fields)
    for row in variants:
        ws2.append([row.get(field, "") for field in variant_fields])

    ws3 = workbook.create_sheet("Needs Review")
    ws3.append(product_fields)
    for row in products:
        if row.get("needs_review") == "YES":
            ws3.append([row.get(field, "") for field in product_fields])

    ws4 = workbook.create_sheet("Excluded Stickers")
    excluded_fields = sorted({key for row in excluded for key in row.keys()}) if excluded else ["source_file", "source_page", "reason"]
    ws4.append(excluded_fields)
    for row in excluded:
        ws4.append([row.get(field, "") for field in excluded_fields])

    try:
        workbook.save(path)
        return path
    except PermissionError:
        alt = path.with_name(f"{path.stem}-{stable_id(len(products), len(variants))}{path.suffix}")
        workbook.save(alt)
        return alt


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    PUBLIC_IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    raw_rows = []
    excluded = []
    unreadable = []
    for default_supplier, path in FILES:
        if not path.exists():
            unreadable.append({"source_file": str(path), "reason": "missing file"})
            continue
        doc = fitz.open(path)
        for page_index in range(doc.page_count):
            try:
                raw_rows.extend(page_rows(default_supplier, path, doc, page_index, excluded))
            except Exception as exc:
                unreadable.append({"source_file": path.name, "source_page": page_index + 1, "reason": repr(exc)})

    products, variants = build_preview(raw_rows)
    product_fields = list(products[0].keys()) if products else []
    variant_fields = list(variants[0].keys()) if variants else []
    write_csv(OUT_DIR / "csl-products-preview.csv", products, product_fields)
    write_csv(OUT_DIR / "csl-variants-preview.csv", variants, variant_fields)
    write_csv(
        OUT_DIR / "csl-excluded-stickers.csv",
        excluded,
        sorted({key for row in excluded for key in row.keys()}) if excluded else ["source_file", "source_page", "reason"],
    )
    workbook_path = write_xlsx(OUT_DIR / "csl-upload-preview.xlsx", products, variants, excluded)
    summary = {
        "source_files": len(FILES),
        "raw_variant_rows": len(raw_rows),
        "products": len(products),
        "variants": len(variants),
        "excluded_sticker_rows_or_pages": len(excluded),
        "unreadable_pages_or_files": unreadable,
        "products_missing_images": sum(1 for product in products if not clean(product.get("main_image_url"))),
        "products_needing_review": sum(1 for product in products if product.get("needs_review") == "YES"),
        "pricing_rule": "PDF price treated as supplier net cost; retail=ceil(cost*1.20); 6pcs_plus=ceil(cost*1.12)",
        "files": {
            "products_csv": str(OUT_DIR / "csl-products-preview.csv"),
            "variants_csv": str(OUT_DIR / "csl-variants-preview.csv"),
            "excluded_csv": str(OUT_DIR / "csl-excluded-stickers.csv"),
            "workbook": str(workbook_path),
            "summary_json": str(OUT_DIR / "csl-preview-summary.json"),
        },
    }
    (OUT_DIR / "csl-preview-summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(summary, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
