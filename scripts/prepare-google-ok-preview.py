import csv
import hashlib
import html
import json
import math
import re
import time
import urllib.parse
import urllib.request
from html.parser import HTMLParser
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "tmp-imports" / "google-ok"
HTML_DIR = OUT_DIR / "sheet-html"
CSV_DIR = OUT_DIR / "sheet-csv"
PREVIEW_DIR = OUT_DIR / "preview"
SPREADSHEET_ID = "1QHDMJGldFz3xAc7MgEJ-moqW-V_P6a11mmv2O3IjxNc"

SHEETS = [
    (1531550958, "MSM"),
    (729656869, "HACHI"),
    (977589521, "LOCAL ITEM"),
    (1494135237, "REPAIR KIT"),
    (2006721384, "SALE ITEMS 1"),
    (560200285, "SALE ITEMS 2"),
    (1931309384, "DECORATIVE LIGHT"),
    (1030173525, "HEADLIGHT LED"),
    (733873277, "MIRROR"),
    (297829640, "MODULE LIGHT"),
    (1728823088, "PARK LIGHT"),
    (731695309, "HEADLIGHT LENS"),
    (431637054, "AXLE CAP"),
    (310375100, "BEARING"),
    (2091921036, "HUB"),
    (1037261863, "TIRE CAP"),
    (1932060014, "MUGS"),
    (1456296230, "  TIRE & WHEEL"),
    (264225889, "RIM"),
    (2032821624, "SOCKET"),
    (575120345, "HANDLE SWITCH"),
    (527396691, "HALO SWITCH"),
    (2126584458, "NEUTRAL"),
    (2067105039, "DOMINO SWITCH"),
    (1612711409, "SWITCH"),
    (321515496, " SEAT & COVER"),
    (574530465, "SEAT ASSY"),
    (1714152982, "SCREW&BOLT"),
    (1960989045, "OTHERS"),
    (1921062297, "RIDING GEAR"),
    (1120536050, "POWER PARTS"),
    (1183712272, "SHAFT"),
    (1148314524, "CLUTCH"),
    (608527999, "ROCKER ARM"),
    (1119182965, "PULLEY"),
    (244530982, "FUEL TANK"),
    (531994058, "PISTON"),
    (1366073817, "FILTER"),
    (175617813, "CARBURETOR"),
    (906406757, "SPARK PLUG"),
    (1849663029, "CABLE"),
    (1331607883, "CAMSHAFT"),
    (1120185491, "MANIFOLD"),
    (1737328708, "GASKET&SEAL"),
    (1727611983, "FOOTREST & PEDAL"),
    (1924905480, "FOOTREST"),
    (1032365199, "EXHAUST"),
    (1724326075, "ELECTRICAL & BATTERY"),
    (1264938005, "GAUGE"),
    (1117033305, "CDI"),
    (414099867, "CHAIN & SPROCKET"),
    (648243160, "BRAKE HOSE"),
    (1787807347, "BRAKE PUMP"),
    (1835393259, "DISC"),
    (1057258732, "CALIPER"),
    (1310468201, "BRAKE SHOE"),
    (11302834, "BRAKE ROD"),
    (431040162, "BRAKE PAD"),
    (1381868843, "BRAKE LEVER"),
    (9256757, "BRAKE ARM & BRAKE CAM"),
    (1837613106, "ASST"),
    (1961906717, "VISOR"),
    (2104337297, "FENDER"),
    (707625481, "BODYCOVER"),
    (1347698804, "OTHER ACCESSORIES"),
    (722924212, "BAR END"),
    (303337954, "STAINLESS STEEL"),
    (1667545294, "STICKER"),
    (467744846, "SPRING"),
    (364045247, "MATTING"),
    (846122485, "HOOK & HOLDER"),
    (291675017, "CLEANING"),
    (1463138727, "LIQUID & SPRAY"),
    (1620148863, "MDL"),
    (89020484, "WINKER LAMP"),
    (905977997, "Sheet65"),
    (2138946243, "TAIL LIGHT LENS"),
    (128283783, "BULB"),
    (259030317, "SIGNAL LIGHT LED"),
    (1689647704, "TAIL LIGHT LED"),
    (349298349, "AUDIO"),
    (419758392, "HORN"),
    (668447376, "HANDLE EXTENSION"),
    (602513165, "HANDLE BAR"),
    (1528961807, "HANDLE GRIP"),
    (1080653914, "STORAGE & SECURITY"),
    (1676102027, "SHOCK"),
    (1216183256, "SHOCK COMPONENT"),
    (1639154816, "BRACKET"),
    (596464665, "."),
]

ATTRIBUTE_HEADERS = {
    "COLOR",
    "COLOUR",
    "MODEL",
    "SIZE",
    "TYPE",
    "FITMENT",
    "FIT",
    "PCS",
    "DESCRIPTION",
}

NON_PRODUCT_TITLES = {
    "",
    ".",
    "SALE",
    "SALE ITEMS",
    "SALE ITEMS 1",
    "SALE ITEMS 2",
    "SALE 促销",
    "CODE",
    "PICTURE",
    "COLOR",
    "MODEL",
    "PRICE",
    "QTY/BOX",
}

CATEGORY_GROUPS = {
    "DECORATIVE LIGHT": "Lights & Electrical",
    "HEADLIGHT LED": "Lights & Electrical",
    "MODULE LIGHT": "Lights & Electrical",
    "PARK LIGHT": "Lights & Electrical",
    "HEADLIGHT LENS": "Lights & Electrical",
    "WINKER LAMP": "Lights & Electrical",
    "TAIL LIGHT LENS": "Lights & Electrical",
    "BULB": "Lights & Electrical",
    "SIGNAL LIGHT LED": "Lights & Electrical",
    "TAIL LIGHT LED": "Lights & Electrical",
    "AUDIO": "Lights & Electrical",
    "HORN": "Lights & Electrical",
    "ELECTRICAL & BATTERY": "Lights & Electrical",
    "CDI": "Lights & Electrical",
    "GAUGE": "Lights & Electrical",
    "MIRROR": "Motorcycle Accessories",
    "AXLE CAP": "Motorcycle Accessories",
    "TIRE CAP": "Motorcycle Accessories",
    "MUGS": "Motorcycle Accessories",
    "SOCKET": "Motorcycle Accessories",
    "HANDLE SWITCH": "Motorcycle Accessories",
    "HALO SWITCH": "Motorcycle Accessories",
    "NEUTRAL": "Motorcycle Accessories",
    "DOMINO SWITCH": "Motorcycle Accessories",
    "SWITCH": "Motorcycle Accessories",
    "RIDING GEAR": "Motorcycle Accessories",
    "VISOR": "Motorcycle Accessories",
    "OTHER ACCESSORIES": "Motorcycle Accessories",
    "BAR END": "Motorcycle Accessories",
    "STAINLESS STEEL": "Motorcycle Accessories",
    "STICKER": "Motorcycle Accessories",
    "MATTING": "Motorcycle Accessories",
    "HOOK & HOLDER": "Motorcycle Accessories",
    "MDL": "Motorcycle Accessories",
    "HANDLE EXTENSION": "Motorcycle Accessories",
    "HANDLE BAR": "Motorcycle Accessories",
    "HANDLE GRIP": "Motorcycle Accessories",
    "STORAGE & SECURITY": "Motorcycle Accessories",
    "BRACKET": "Motorcycle Accessories",
    "SEAT & COVER": "Body & Styling",
    "SEAT ASSY": "Body & Styling",
    "FENDER": "Body & Styling",
    "BODYCOVER": "Body & Styling",
    "FUEL TANK": "Body & Styling",
    "TIRE & WHEEL": "Wheels & Tires",
    "RIM": "Wheels & Tires",
    "HUB": "Wheels & Tires",
    "BEARING": "Wheels & Tires",
    "CHAIN & SPROCKET": "Drive Train",
    "SHAFT": "Drive Train",
    "CLUTCH": "Drive Train",
    "PULLEY": "Drive Train",
    "FOOTREST & PEDAL": "Controls & Footrests",
    "FOOTREST": "Controls & Footrests",
    "BRAKE HOSE": "Brake System",
    "BRAKE PUMP": "Brake System",
    "DISC": "Brake System",
    "CALIPER": "Brake System",
    "BRAKE SHOE": "Brake System",
    "BRAKE ROD": "Brake System",
    "BRAKE PAD": "Brake System",
    "BRAKE LEVER": "Brake System",
    "BRAKE ARM & BRAKE CAM": "Brake System",
    "ROCKER ARM": "Engine Parts",
    "PISTON": "Engine Parts",
    "FILTER": "Engine Parts",
    "CARBURETOR": "Engine Parts",
    "SPARK PLUG": "Engine Parts",
    "CABLE": "Engine Parts",
    "CAMSHAFT": "Engine Parts",
    "MANIFOLD": "Engine Parts",
    "GASKET&SEAL": "Engine Parts",
    "EXHAUST": "Engine Parts",
    "POWER PARTS": "Engine Parts",
    "REPAIR KIT": "Engine Parts",
    "SHOCK": "Suspension",
    "SHOCK COMPONENT": "Suspension",
    "CLEANING": "Care & Maintenance",
    "LIQUID & SPRAY": "Care & Maintenance",
}


def clean(value):
    if value is None:
        return ""
    text = html.unescape(str(value)).replace("\xa0", " ").strip()
    return re.sub(r"\s+", " ", text)


def slugify(value, fallback="item"):
    text = clean(value).lower()
    text = re.sub(r"[^a-z0-9]+", "-", text)
    return re.sub(r"-+", "-", text).strip("-") or fallback


def stable_id(*parts):
    raw = "|".join(clean(part) for part in parts)
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:8]


def parse_number(value):
    text = clean(value).replace(",", "")
    if not text:
        return None
    match = re.search(r"\d+(?:\.\d+)?", text)
    if not match:
        return None
    return float(match.group(0))


def english_frontend_text(value):
    text = re.sub(r"[^\x00-\x7F]+", " ", clean(value))
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"\(\s*\)", "", text)
    text = re.sub(r"\s+([),./-])", r"\1", text)
    text = re.sub(r"([(])\s+", r"\1", text)
    return text.strip(" -_/.,")


def title_keep_codes(value):
    value = english_frontend_text(value)
    words = []
    for part in clean(value).split(" "):
        if re.search(r"\d", part) or part.isupper() or "-" in part or "/" in part or "&" in part:
            words.append(part.upper())
        else:
            words.append(part.capitalize())
    return " ".join(words)


def sheet_display_name(sheet_name):
    return clean(sheet_name).strip(".").strip() or "OK Catalog"


def category_for_sheet(sheet_name):
    child = sheet_display_name(sheet_name)
    group = CATEGORY_GROUPS.get(child.upper(), "OK Catalog")
    return f"Motorcycle Parts > {group} > {title_keep_codes(child)}"


def product_name(source_title, sheet_name):
    title = clean(source_title) or sheet_display_name(sheet_name)
    name = title_keep_codes(title)
    if not name.upper().startswith("OK "):
        name = f"OK {name}"
    return name


def looks_like_code(value):
    text = clean(value).upper()
    if not text or text in NON_PRODUCT_TITLES:
        return False
    return text.startswith("OK")


def is_title_row(texts):
    values = [clean(text) for text in texts if clean(text)]
    if not values:
        return False
    joined = " ".join(values).upper()
    if is_supplier_note_text(joined):
        return False
    if any(token in joined for token in ["CODE", "PRICE", "QTY/BOX", "PICTURE"]):
        return False
    if looks_like_code(values[0]):
        return False
    if any(re.fullmatch(r"[\d,.]+", value) for value in values):
        return False
    if joined in NON_PRODUCT_TITLES:
        return False
    if not re.search(r"[A-Z]", english_frontend_text(joined).upper()):
        return False
    return bool(re.search(r"[A-Z]", joined))


def is_supplier_note_text(value):
    text = clean(value).upper()
    if not text:
        return False
    note_tokens = [
        "不需要",
        "价格",
        "變動",
        "变动",
        "请注意",
        "謝謝",
        "谢谢",
        "各种型号",
        "合 计",
        "合计",
        "再便宜",
        "原来",
        "点数",
        "點數",
        "开始",
        "下调",
        "下調",
        "促销",
        "促銷",
        "箱",
    ]
    if any(token in text for token in note_tokens):
        return True
    return False


class SheetHtmlParser(HTMLParser):
    def __init__(self, gid):
        super().__init__(convert_charrefs=True)
        self.gid = str(gid)
        self.rows = []
        self.in_tr = False
        self.row_id = None
        self.current_row = []
        self.in_cell = False
        self.current_cell = None
        self.cell_tag = None
        self.data_parts = []

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "tr":
            self.in_tr = True
            self.row_id = None
            self.current_row = []
        elif self.in_tr and tag == "th":
            row_match = re.match(rf"{re.escape(self.gid)}R(\d+)$", attrs.get("id", ""))
            if row_match:
                self.row_id = int(row_match.group(1)) + 1
        elif self.in_tr and tag == "td":
            self.in_cell = True
            self.cell_tag = "td"
            self.data_parts = []
            colspan = int(attrs.get("colspan", "1")) if attrs.get("colspan", "1").isdigit() else 1
            self.current_cell = {"text": "", "images": [], "colspan": max(1, colspan)}
        elif self.in_cell and tag == "img" and self.current_cell is not None:
            src = attrs.get("src")
            if src:
                self.current_cell["images"].append(html.unescape(src))

    def handle_data(self, data):
        if self.in_cell:
            self.data_parts.append(data)

    def handle_endtag(self, tag):
        if tag == "td" and self.in_cell and self.current_cell is not None:
            self.current_cell["text"] = clean(" ".join(self.data_parts))
            colspan = self.current_cell.pop("colspan", 1)
            self.current_row.append(self.current_cell)
            for _ in range(colspan - 1):
                self.current_row.append({"text": "", "images": []})
            self.current_cell = None
            self.in_cell = False
        elif tag == "tr" and self.in_tr:
            if self.row_id is not None:
                self.rows.append({"row": self.row_id, "cells": self.current_row})
            self.in_tr = False


def fetch_sheet_html(gid, sheet_name):
    HTML_DIR.mkdir(parents=True, exist_ok=True)
    path = HTML_DIR / f"{str(gid)}-{slugify(sheet_name)}.html"
    if path.exists() and path.stat().st_size > 1000:
        return path.read_text(encoding="utf-8", errors="ignore")
    url = f"https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/edit?gid={gid}"
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(request, timeout=90) as response:
        text = response.read().decode("utf-8", errors="ignore")
    path.write_text(text, encoding="utf-8")
    time.sleep(0.25)
    return text


def parse_sheet_rows(gid, sheet_name):
    text = fetch_sheet_html(gid, sheet_name)
    parser = SheetHtmlParser(gid)
    parser.feed(text)
    return parser.rows


def fetch_sheet_csv(sheet_name):
    CSV_DIR.mkdir(parents=True, exist_ok=True)
    path = CSV_DIR / f"{slugify(sheet_name)}.csv"
    if path.exists() and path.stat().st_size > 10:
        return path
    encoded_sheet = urllib.parse.quote(sheet_name, safe="")
    url = f"https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/gviz/tq?tqx=out:csv&headers=0&sheet={encoded_sheet}"
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(request, timeout=90) as response:
        data = response.read()
    path.write_bytes(data)
    time.sleep(0.15)
    return path


def parse_sheet_csv_rows(sheet_name):
    path = fetch_sheet_csv(sheet_name)
    rows = []
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        for values in csv.reader(fh):
            rows.append(values)
    return rows


def header_map(cells):
    mapping = {}
    for index, cell in enumerate(cells):
        text = clean(cell["text"]).upper()
        compact_text = re.sub(r"\s+", "", text)
        if text in {"CODE", "ITEM CODE", "SKU"}:
            mapping["code"] = index
        elif text in {
            "COLOR",
            "COLOUR",
            "MODEL",
            "SIZE",
            "TYPE",
            "FITMENT",
            "FIT",
            "DESCRIPTION",
            "PCS",
            "BRAND",
            "SPECIFICATION",
        }:
            mapping.setdefault("attrs", []).append((index, text))
        elif compact_text in {"MODEL/SIZE", "MODEL/SIZES", "MODELSIZE", "MODEL/COLOR", "MODELCOLOR"}:
            mapping.setdefault("attrs", []).append((index, "MODEL"))
        elif text in {"QTY/BOX", "QTY", "BOX"}:
            mapping["qty"] = index
        elif text in {"PRICE", "COST", "COST PRICE"}:
            mapping["price"] = index
    return mapping if "code" in mapping else {}


def cell_text(cells, index):
    if index is None or index >= len(cells):
        return ""
    return clean(cells[index]["text"])


def cell_images(cells):
    images = []
    for cell in cells:
        images.extend(cell.get("images", []))
    return images


def first_code_and_price(cells):
    code = ""
    price = None
    for cell in cells:
        text = cell["text"] if isinstance(cell, dict) else cell
        if not code and looks_like_code(text):
            code = clean(text)
            continue
        if code:
            number = parse_number(text)
            if number is not None:
                price = number
    return code, price


def infer_qty_and_price(cells, code_col, qty_col=None, price_col=None):
    if price_col is not None:
        price = parse_number(cell_text(cells, price_col))
        qty = cell_text(cells, qty_col) if qty_col is not None else ""
        return qty, price
    numeric_cells = []
    for index, cell in enumerate(cells):
        if index <= code_col:
            continue
        number = parse_number(cell["text"])
        if number is not None:
            numeric_cells.append((index, cell["text"], number))
    if not numeric_cells:
        return "", None
    price = numeric_cells[-1][2]
    qty = numeric_cells[-2][1] if len(numeric_cells) >= 2 else ""
    return clean(qty), price


def derive_csv_row_offset(csv_rows, html_rows):
    html_signatures = []
    for row in html_rows:
        code, price = first_code_and_price(row["cells"])
        if code and price is not None:
            html_signatures.append((code, price, row["row"]))
    if not html_signatures:
        return 1
    for csv_index, values in enumerate(csv_rows[:140]):
        code, price = first_code_and_price(values)
        if not code or price is None:
            continue
        for html_code, html_price, html_row in html_signatures:
            if code == html_code and abs(price - html_price) < 0.001:
                return html_row - csv_index
    return 1


def csv_values_to_cells(values, row_images=None):
    cells = []
    for index, value in enumerate(values):
        cells.append({"text": clean(value), "images": row_images if index == 0 and row_images else []})
    return cells


def parse_products_from_sheet(gid, sheet_name):
    html_rows = parse_sheet_rows(gid, sheet_name)
    images_by_row = {row["row"]: cell_images(row["cells"]) for row in html_rows if cell_images(row["cells"])}
    images_by_code = {}
    for row in html_rows:
        code, _ = first_code_and_price(row["cells"])
        images = cell_images(row["cells"])
        if code and images:
            images_by_code[code] = images
    csv_rows = parse_sheet_csv_rows(sheet_name)
    row_offset = derive_csv_row_offset(csv_rows, html_rows)
    rows = [
        {"row": index + row_offset, "cells": csv_values_to_cells(values, images_by_row.get(index + row_offset, []))}
        for index, values in enumerate(csv_rows)
    ]
    active_header = {}
    current_title = sheet_display_name(sheet_name)
    last_image_for_title = {}
    variants = []

    for row in rows:
        cells = row["cells"]
        texts = [cell["text"] for cell in cells]
        maybe_header = header_map(cells)
        if maybe_header:
            active_header = maybe_header
            continue

        if is_title_row(texts):
            values = [clean(text) for text in texts if clean(text)]
            title = " ".join(values[:2])
            if title.upper() not in NON_PRODUCT_TITLES:
                current_title = title
            continue

        if not active_header:
            continue

        code = cell_text(cells, active_header.get("code"))
        qty_box, price = infer_qty_and_price(
            cells,
            active_header.get("code", 0),
            active_header.get("qty"),
            active_header.get("price"),
        )
        if not looks_like_code(code) or price is None:
            continue

        attrs = {}
        for index, label in active_header.get("attrs", []):
            value = cell_text(cells, index)
            if value:
                attrs[label.lower().replace("/", "_")] = value
        images = images_by_code.get(code, cell_images(cells))
        title_key = f"{sheet_name}|{current_title}"
        if images:
            last_image_for_title[title_key] = images[0]
        image_url = images[0] if images else last_image_for_title.get(title_key, "")
        variants.append(
            {
                "sheet_id": gid,
                "sheet": sheet_name,
                "row": row["row"],
                "source_title": current_title,
                "code": code,
                "color": attrs.get("color") or attrs.get("colour") or "",
                "model": attrs.get("model")
                or attrs.get("fitment")
                or attrs.get("fit")
                or attrs.get("description")
                or attrs.get("brand")
                or "",
                "size": attrs.get("size") or attrs.get("type") or attrs.get("pcs") or attrs.get("specification") or "",
                "qty_box": qty_box,
                "cost": price,
                "source_image_url": image_url,
            }
        )

    grouped = {}
    for variant in variants:
        key = (sheet_name, variant["source_title"])
        grouped.setdefault(key, []).append(variant)
    return grouped


def variant_label(variant):
    parts = []
    for key in ("color", "model", "size"):
        value = clean(variant.get(key))
        if value and value.upper() not in {"N/A", "NA"}:
            parts.append(value)
    parts.append(variant["code"])
    return " / ".join(parts)


def make_preview():
    products = []
    variants = []
    errors = []
    for index, (gid, sheet_name) in enumerate(SHEETS, start=1):
        try:
            grouped = parse_products_from_sheet(gid, sheet_name)
            print(f"[{index}/{len(SHEETS)}] {sheet_display_name(sheet_name)}: {len(grouped)} products")
        except Exception as exc:
            errors.append({"sheet": sheet_name, "error": str(exc)})
            print(f"[{index}/{len(SHEETS)}] {sheet_display_name(sheet_name)} failed: {exc}")
            continue
        for (_, title), group_variants in grouped.items():
            preview_id = f"ok-{slugify(sheet_name)}-{stable_id(sheet_name, title)}"
            costs = [variant["cost"] for variant in group_variants]
            retail_prices = [math.ceil(cost * 1.20) for cost in costs]
            wholesale_prices = [math.ceil(cost * 1.12) for cost in costs]
            image_url = next((variant["source_image_url"] for variant in group_variants if variant["source_image_url"]), "")
            needs = []
            if not image_url:
                needs.append("missing image")
            if not title:
                needs.append("missing product title")
            category = category_for_sheet(sheet_name)
            products.append(
                {
                    "preview_id": preview_id,
                    "supplier": "OK",
                    "sheet": sheet_display_name(sheet_name),
                    "category": category,
                    "source_title": title,
                    "english_product_name": product_name(title, sheet_name),
                    "variant_count": len(group_variants),
                    "cost_min": min(costs),
                    "cost_max": max(costs),
                    "retail_min_1_5pcs": min(retail_prices),
                    "retail_max_1_5pcs": max(retail_prices),
                    "wholesale_min_6pcs_plus": min(wholesale_prices),
                    "wholesale_max_6pcs_plus": max(wholesale_prices),
                    "moq": 1,
                    "main_image_url": image_url,
                    "needs_review": "YES" if needs else "NO",
                    "review_note": "; ".join(needs),
                }
            )
            seen_variant_keys = set()
            for variant_index, variant in enumerate(group_variants, start=1):
                cost = variant["cost"]
                label = variant_label(variant)
                key = (variant["code"], label, cost)
                if key in seen_variant_keys:
                    continue
                seen_variant_keys.add(key)
                variants.append(
                    {
                        "preview_id": preview_id,
                        "variant_sku": f"OK-{slugify(variant['code']).upper()}-{stable_id(preview_id, label).upper()}"[:96],
                        "source_code": variant["code"],
                        "variant_name": label,
                        "color": variant["color"],
                        "model": variant["model"],
                        "size": variant["size"],
                        "qty_box": variant["qty_box"],
                        "cost": cost,
                        "retail_price_1_5pcs": math.ceil(cost * 1.20),
                        "wholesale_price_6pcs_plus": math.ceil(cost * 1.12),
                        "source_image_url": variant["source_image_url"] or image_url,
                        "source_sheet": sheet_display_name(sheet_name),
                        "source_row": variant["row"],
                        "sort_order": variant_index,
                    }
                )
    return products, variants, errors


def write_csv(path, rows, fields):
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(path, products, variants):
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = "Products Preview"
    product_fields = list(products[0].keys()) if products else []
    ws.append(product_fields)
    for row in products:
        ws.append([row.get(field, "") for field in product_fields])
    ws2 = workbook.create_sheet("Variants Preview")
    variant_fields = list(variants[0].keys()) if variants else []
    ws2.append(variant_fields)
    for row in variants:
        ws2.append([row.get(field, "") for field in variant_fields])
    ws3 = workbook.create_sheet("Missing Images")
    missing = [product for product in products if product["needs_review"] == "YES"]
    ws3.append(product_fields)
    for row in missing:
        ws3.append([row.get(field, "") for field in product_fields])
    workbook.save(path)


def main():
    PREVIEW_DIR.mkdir(parents=True, exist_ok=True)
    products, variants, errors = make_preview()
    if products:
        write_csv(PREVIEW_DIR / "ok-products-preview.csv", products, list(products[0].keys()))
    if variants:
        write_csv(PREVIEW_DIR / "ok-variants-preview.csv", variants, list(variants[0].keys()))
    write_xlsx(PREVIEW_DIR / "ok-upload-preview.xlsx", products, variants)
    summary = {
        "products": len(products),
        "variants": len(variants),
        "sheets": len(SHEETS),
        "sheets_with_errors": len(errors),
        "needs_review_products": sum(1 for product in products if product["needs_review"] == "YES"),
        "pricing_rule": "cost=PRICE from supplier sheet, retail=ceil(cost*1.20), 6pcs_plus=ceil(cost*1.12)",
        "by_sheet": {},
        "errors": errors,
        "files": {
            "products_csv": str(PREVIEW_DIR / "ok-products-preview.csv"),
            "variants_csv": str(PREVIEW_DIR / "ok-variants-preview.csv"),
            "workbook": str(PREVIEW_DIR / "ok-upload-preview.xlsx"),
            "summary_json": str(PREVIEW_DIR / "ok-preview-summary.json"),
        },
    }
    for product in products:
        summary["by_sheet"][product["sheet"]] = summary["by_sheet"].get(product["sheet"], 0) + 1
    (PREVIEW_DIR / "ok-preview-summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(summary, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
