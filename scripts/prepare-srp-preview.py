import csv
import hashlib
import json
import math
import os
import re
from pathlib import Path

import openpyxl
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "tmp-imports" / "srp-xlsx"
OUT_DIR = ROOT / "tmp-imports" / "srp-preview"
PUBLIC_IMAGE_DIR = ROOT / "public" / "product-imports" / "srp-preview"

FILES = [
    ("gille", SOURCE_DIR / "gille-updated-stocks.xlsx"),
    ("zebra", SOURCE_DIR / "zebra-srp-updated-stocks.xlsx"),
]

SHEET_CATEGORY = {
    "HALF FACE": "Motorcycle Parts > Helmets > Half Face Helmets",
    "FULL FACE": "Motorcycle Parts > Helmets > Full Face Helmets",
    "DUAL SPORTS": "Motorcycle Parts > Helmets > Dual Sport Helmets",
    "MODULAR": "Motorcycle Parts > Helmets > Modular Helmets",
    "MOTOCROSS": "Motorcycle Parts > Helmets > Motocross Helmets",
}

BAD_VALUES = {
    "ITEM NAME",
    "COLOR",
    "COLOUR",
    "SIZE",
    "AVAILABLE",
    "SPECIFICATION",
    "NEW PRICE",
    "SRP NEW PRICE",
    "SRP PRICE",
    "PRICE",
    "SRP",
}

SIZE_NAMES = {"XS", "S", "M", "L", "XL", "XXL", "2XL", "3XL", "4XL", "5XL", "XXXL"}


def clean(value):
    if value is None:
        return ""
    text = str(value).replace("\t", " ").strip()
    return re.sub(r"\s+", " ", text)


def slugify(value, fallback="item"):
    text = clean(value).lower()
    text = re.sub(r"[^a-z0-9]+", "-", text)
    return re.sub(r"-+", "-", text).strip("-") or fallback


def stable_id(*parts):
    raw = "|".join(clean(part) for part in parts)
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:8]


def normalize_size(value):
    text = clean(value).upper().replace(" ", "")
    aliases = {"2XL": "XXL", "3XL": "XXXL"}
    return aliases.get(text, text)


def is_size(value):
    return normalize_size(value) in SIZE_NAMES


def is_label(value):
    upper = clean(value).upper()
    return not upper or upper in BAD_VALUES or "PRICE" in upper or upper.startswith("1*")


def is_bad_color(value):
    upper = clean(value).upper()
    if is_label(upper):
        return True
    if re.fullmatch(r"[\d.,]+", upper):
        return True
    if re.fullmatch(r"[SMLX0-9/ -]+", upper):
        return True
    if upper in {"YES", "NO", "N/A", "NA", "AVAILABLE"}:
        return True
    return False


def normalize_brand(name):
    upper = name.upper()
    if "GILLE" in upper:
        return "GILLE"
    if "ZEBRA" in upper:
        return "ZEBRA"
    return ""


def title_keep_codes(text):
    words = []
    for part in clean(text).split(" "):
        if re.search(r"\d", part) or part.isupper() or "-" in part or "/" in part:
            words.append(part.upper())
        else:
            words.append(part.capitalize())
    return " ".join(words)


def polish_name(raw_name, sheet_name):
    name = clean(raw_name)
    name = re.sub(r"\bW/D\b", "Graphic", name, flags=re.I)
    name = re.sub(r"\bWD\b", "Graphic", name, flags=re.I)
    name = re.sub(r"\(\s*REVO\s*VISOR\s*\)", " Revo Visor", name, flags=re.I)
    name = re.sub(r"\bPLAIN\b", "Plain", name, flags=re.I)
    name = re.sub(r"\s+", " ", name)
    base = title_keep_codes(name)
    suffix_by_sheet = {
        "HALF FACE": "Half Face Helmet",
        "FULL FACE": "Full Face Helmet",
        "DUAL SPORTS": "Dual Sport Helmet",
        "MODULAR": "Modular Helmet",
        "MOTOCROSS": "Motocross Helmet",
    }
    suffix = suffix_by_sheet.get(sheet_name.upper(), "Helmet")
    if suffix.upper() not in base.upper():
        base = f"{base} {suffix}"
    return base


def dedupe_product_names(products):
    groups = {}
    for product in products:
        groups.setdefault(product["english_product_name"], []).append(product)
    for _, group in groups.items():
        if len(group) <= 1:
            continue
        for idx, product in enumerate(group, start=1):
            product["english_product_name"] = f"{product['english_product_name']} Design {idx}"


def sheet_header_blocks(ws):
    headers = []
    for row in ws.iter_rows():
        for cell in row:
            if clean(cell.value).upper() == "ITEM NAME":
                headers.append((cell.row, cell.column))
    headers.sort()
    blocks = []
    for idx, (row, col) in enumerate(headers):
        same_row_next = [next_col for next_row, next_col in headers[idx + 1 :] if next_row == row and next_col > col]
        end_col = min(same_row_next) - 1 if same_row_next else min(ws.max_column, col + 8)
        blocks.append((row, col, end_col))
    return blocks


def read_price(ws, header_row, item_col):
    for rr in range(header_row + 1, min(ws.max_row, header_row + 30) + 1):
        value = ws.cell(rr, item_col).value
        text = clean(value).upper()
        if "PRICE" not in text and text != "SRP":
            continue
        for pr in range(rr + 1, min(ws.max_row, rr + 8) + 1):
            raw = ws.cell(pr, item_col).value
            if isinstance(raw, (int, float)) and raw > 0:
                return float(raw)
            raw_text = clean(raw).replace(",", "")
            if re.fullmatch(r"\d+(?:\.\d+)?", raw_text):
                return float(raw_text)
    for rr in range(header_row + 1, min(ws.max_row, header_row + 30) + 1):
        raw = ws.cell(rr, item_col).value
        if isinstance(raw, (int, float)) and raw >= 100:
            return float(raw)
    return None


def read_name(ws, header_row, item_col):
    for rr in range(header_row + 1, min(ws.max_row, header_row + 20) + 1):
        text = clean(ws.cell(rr, item_col).value)
        if not text or is_label(text):
            continue
        if re.fullmatch(r"\d+(?:\.\d+)?", text):
            continue
        return text
    return ""


def read_spec(ws, header_row, item_col):
    for rr in range(header_row + 1, min(ws.max_row, header_row + 20) + 1):
        if clean(ws.cell(rr, item_col).value).upper() == "SPECIFICATION":
            return clean(ws.cell(rr + 1, item_col).value)
    return ""


def read_sizes(ws, header_row, start_col, end_col):
    sizes = []
    for cc in range(start_col + 2, end_col + 1):
        for rr in range(header_row, min(ws.max_row, header_row + 4) + 1):
            value = clean(ws.cell(rr, cc).value)
            if is_size(value):
                sizes.append((cc, normalize_size(value)))
                break
    return sizes


def has_stock_mark(value):
    text = clean(value).upper()
    if not text:
        return False
    if text in {"0", "-", "NO", "N/A", "NA"}:
        return False
    return True


def read_variants(ws, header_row, start_col, end_col):
    color_col = start_col + 1
    size_cols = read_sizes(ws, header_row, start_col, end_col)
    variants = []
    seen = set()
    for rr in range(header_row + 1, min(ws.max_row, header_row + 45) + 1):
        color = clean(ws.cell(rr, color_col).value)
        if is_bad_color(color):
            continue
        if size_cols:
            active_sizes = [size for cc, size in size_cols if has_stock_mark(ws.cell(rr, cc).value)]
            if not active_sizes:
                active_sizes = [size for _, size in size_cols]
            for size in active_sizes:
                key = (color.upper(), size.upper())
                if key not in seen:
                    variants.append({"color": color, "size": size})
                    seen.add(key)
        else:
            key = (color.upper(), "")
            if key not in seen:
                variants.append({"color": color, "size": ""})
                seen.add(key)
    return variants


def image_anchor(image):
    anchor = getattr(image, "anchor", None)
    marker = getattr(anchor, "_from", None)
    if not marker:
        return None
    return marker.row + 1, marker.col + 1


def extract_sheet_images(wb, ws, source_slug):
    images = []
    image_dir = PUBLIC_IMAGE_DIR / source_slug / slugify(ws.title)
    image_dir.mkdir(parents=True, exist_ok=True)
    for idx, image in enumerate(getattr(ws, "_images", []), start=1):
        anchor = image_anchor(image)
        if not anchor:
            continue
        ext = "png"
        try:
            raw = image._data()
        except Exception:
            continue
        filename = f"{idx:03d}-{anchor[0]}-{anchor[1]}.{ext}"
        path = image_dir / filename
        path.write_bytes(raw)
        images.append({"row": anchor[0], "col": anchor[1], "path": path})
    images.sort(key=lambda item: (item["row"], item["col"]))
    return images


def choose_image(images, header_row, start_col, end_col):
    candidates = [
        image
        for image in images
        if image["col"] >= start_col and image["col"] <= max(end_col, start_col + 2) and image["row"] <= header_row
    ]
    if not candidates:
        candidates = [image for image in images if image["col"] >= start_col and image["col"] <= end_col + 2]
    if not candidates:
        return ""
    image = sorted(candidates, key=lambda item: (abs(item["col"] - start_col), item["row"]))[0]
    return "/" + str(image["path"].relative_to(ROOT / "public")).replace("\\", "/")


def process_workbook(source_slug, path):
    wb = openpyxl.load_workbook(path)
    products = []
    variants = []
    for ws in wb.worksheets:
        if ws.title.upper() not in SHEET_CATEGORY:
            continue
        images = extract_sheet_images(wb, ws, source_slug)
        for header_row, start_col, end_col in sheet_header_blocks(ws):
            raw_name = read_name(ws, header_row, start_col)
            srp = read_price(ws, header_row, start_col)
            if not raw_name or not srp:
                continue
            product_id = f"{source_slug}-{slugify(ws.title)}-{start_col}-{stable_id(raw_name, srp)}"
            cost = math.ceil(srp * 0.30)
            retail = math.ceil(cost * 1.20)
            wholesale_6 = math.ceil(cost * 1.12)
            product_variants = read_variants(ws, header_row, start_col, end_col)
            image_path = choose_image(images, header_row, start_col, end_col)
            product = {
                "preview_id": product_id,
                "source_file": path.name,
                "sheet": ws.title,
                "category": SHEET_CATEGORY[ws.title.upper()],
                "brand": normalize_brand(raw_name),
                "supplier_name": raw_name,
                "english_product_name": polish_name(raw_name, ws.title),
                "specification": read_spec(ws, header_row, start_col),
                "srp": round(srp, 2),
                "computed_cost_30pct": cost,
                "retail_price_1_5pcs": retail,
                "wholesale_price_6pcs_plus": wholesale_6,
                "moq": 1,
                "main_image": image_path,
                "variant_count": len(product_variants),
                "needs_review": "NO" if image_path and product_variants else "YES",
                "review_note": "" if image_path and product_variants else "Missing image or color/size variant data",
            }
            products.append(product)
            if not product_variants:
                product_variants = [{"color": "", "size": ""}]
            for idx, variant in enumerate(product_variants, start=1):
                variants.append(
                    {
                        "preview_id": product_id,
                        "variant_sku": f"{slugify(product['brand'] or source_slug)}-{stable_id(product_id, variant['color'], variant['size'])}",
                        "color": variant["color"],
                        "size": variant["size"],
                        "image": image_path,
                    }
                )
    return products, variants


def write_csv(path, rows, fieldnames):
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(path, products, variants):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products Preview"
    product_fields = list(products[0].keys()) if products else []
    ws.append(product_fields)
    for row in products:
        ws.append([row.get(field, "") for field in product_fields])
    ws2 = wb.create_sheet("Variants Preview")
    variant_fields = list(variants[0].keys()) if variants else []
    ws2.append(variant_fields)
    for row in variants:
        ws2.append([row.get(field, "") for field in variant_fields])
    wb.save(path)


def make_contact_sheet(products):
    thumbs = []
    for product in products:
        image_rel = product["main_image"].lstrip("/")
        image_path = ROOT / "public" / image_rel
        if not image_path.exists():
            continue
        try:
            img = Image.open(image_path).convert("RGB")
        except Exception:
            continue
        img.thumbnail((210, 160))
        thumbs.append((product, img.copy()))
    if not thumbs:
        return None
    width = 1200
    card_w = 300
    card_h = 260
    cols = max(1, width // card_w)
    rows = math.ceil(len(thumbs) / cols)
    sheet = Image.new("RGB", (width, rows * card_h), "white")
    draw = ImageDraw.Draw(sheet)
    font = ImageFont.load_default()
    for idx, (product, img) in enumerate(thumbs):
        x = (idx % cols) * card_w
        y = (idx // cols) * card_h
        sheet.paste(img, (x + (card_w - img.width) // 2, y + 8))
        lines = [
            product["english_product_name"][:38],
            f"SRP {product['srp']} | Cost {product['computed_cost_30pct']}",
            f"Retail {product['retail_price_1_5pcs']} | 6+ {product['wholesale_price_6pcs_plus']}",
            product["sheet"],
        ]
        ty = y + 175
        for line in lines:
            draw.text((x + 10, ty), line, fill=(20, 20, 20), font=font)
            ty += 18
        draw.rectangle((x, y, x + card_w - 1, y + card_h - 1), outline=(220, 220, 220))
    output = OUT_DIR / "srp-contact-sheet.jpg"
    sheet.save(output, "JPEG", quality=88)
    return output


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    PUBLIC_IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    all_products = []
    all_variants = []
    for source_slug, path in FILES:
        products, variants = process_workbook(source_slug, path)
        all_products.extend(products)
        all_variants.extend(variants)
    dedupe_product_names(all_products)

    if all_products:
        write_csv(OUT_DIR / "srp-products-preview.csv", all_products, list(all_products[0].keys()))
    if all_variants:
        write_csv(OUT_DIR / "srp-variants-preview.csv", all_variants, list(all_variants[0].keys()))
    write_xlsx(OUT_DIR / "srp-upload-preview.xlsx", all_products, all_variants)
    contact = make_contact_sheet(all_products)

    summary = {
        "products": len(all_products),
        "variants": len(all_variants),
        "needs_review_products": sum(1 for product in all_products if product["needs_review"] == "YES"),
        "by_source": {},
        "by_sheet": {},
        "pricing_rule": "cost=ceil(SRP*0.30), retail=ceil(cost*1.20), 6pcs_plus=ceil(cost*1.12)",
        "files": {
            "products_csv": str(OUT_DIR / "srp-products-preview.csv"),
            "variants_csv": str(OUT_DIR / "srp-variants-preview.csv"),
            "workbook": str(OUT_DIR / "srp-upload-preview.xlsx"),
            "contact_sheet": str(contact) if contact else "",
        },
    }
    for product in all_products:
        summary["by_source"][product["source_file"]] = summary["by_source"].get(product["source_file"], 0) + 1
        key = f"{product['source_file']} / {product['sheet']}"
        summary["by_sheet"][key] = summary["by_sheet"].get(key, 0) + 1
    (OUT_DIR / "srp-preview-summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
