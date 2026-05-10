import hashlib
import json
import math
import os
import posixpath
import re
import sys
import time
import urllib.parse
import urllib.request
import zipfile
from io import BytesIO
from xml.etree import ElementTree as ET

import openpyxl
from PIL import Image


BASE_DIR = r"C:\Users\Administrator\Desktop\二维码"

FILES = [
    "GILLE UPDATED STOCKS(1).xlsx",
    "ZEBRA SRP UPDATED STOCKS(1).xlsx",
    "ZEBRA HALF FACE, FULL FACE & MODULAR UPDATED STOCKS(1).xlsx",
    "HNJ FULL FACE UPDATED STOCKS(1).xlsx",
    "HNJ HALF FACE HELMETS UPDATED STOCKS(1).xlsx",
    "MOB, CAP TYPE & KIDS HELMET UPDATED STOCKS(1).xlsx",
    "ACCESSORIES AVAILABLE STOCKS(1).xlsx",
    "MOTOBOX AVAILABLE STOCKS(1).xlsx",
    "HNJ MODULAR UPDATED STOCKS(1).xlsx",
]

BAD_VALUES = {
    "ITEM NAME",
    "COLOR",
    "COLOUR",
    "SIZE",
    "AVAILABLE",
    "SPECIFICATION",
    "NEW PRICE",
    "SRP NEW PRICE",
    "SRP PRICE",
    "PRICE",
}


def load_env():
    env_path = os.path.join(os.getcwd(), ".env.local")
    values = {}
    with open(env_path, "r", encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            values[key.strip()] = value.strip().strip('"').strip("'")
    url = values.get("NEXT_PUBLIC_SUPABASE_URL")
    service_key = values.get("SUPABASE_SERVICE_ROLE_KEY")
    if not url or not service_key:
        raise RuntimeError("Missing Supabase server credentials in .env.local")
    return url.rstrip("/"), service_key


SUPABASE_URL, SERVICE_KEY = load_env()


def request_json(method, table, query="", payload=None, prefer="return=representation"):
    url = f"{SUPABASE_URL}/rest/v1/{table}{query}"
    data = None if payload is None else json.dumps(payload).encode("utf-8")
    headers = {
        "apikey": SERVICE_KEY,
        "Authorization": f"Bearer {SERVICE_KEY}",
        "Content-Type": "application/json",
    }
    if prefer:
        headers["Prefer"] = prefer
    req = urllib.request.Request(url, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            body = resp.read().decode("utf-8")
            if not body:
                return None
            return json.loads(body)
    except urllib.error.HTTPError as exc:
        details = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"{method} {table} failed: {exc.code} {details}") from exc


def clean(value):
    if value is None:
        return ""
    text = str(value).replace("\t", " ").strip()
    return re.sub(r"\s+", " ", text)


def slugify(text, fallback="item"):
    text = clean(text).lower()
    text = re.sub(r"[^a-z0-9]+", "-", text)
    text = re.sub(r"-+", "-", text).strip("-")
    return text or fallback


def stable_suffix(*parts):
    raw = "|".join(str(part) for part in parts)
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:8]


def is_label(text):
    upper = clean(text).upper()
    return upper in BAD_VALUES or "PRICE" in upper or upper.startswith("1*") or upper == "SPECIFICATION"


def is_bad_color(text):
    upper = clean(text).upper()
    return (
        not upper
        or upper in BAD_VALUES
        or "ĄĖ" in upper
        or re.fullmatch(r"[SMLX0-9\- /]+", upper) is not None
        or upper.startswith("1*")
        or "PRICE" in upper
    )


def first_price_from_line(text):
    text = clean(text)
    match = re.search(r"PCS\s*-\s*([\d,]+(?:\.\d+)?)\s*P\b", text, flags=re.I)
    if not match:
        match = re.search(r"(?:^|\s)-\s*([\d,]+(?:\.\d+)?)\s*P\b", text, flags=re.I)
    if match:
        return float(match.group(1).replace(",", ""))
    numbers = [float(x.replace(",", "")) for x in re.findall(r"\d[\d,]*(?:\.\d+)?", text)]
    return max(numbers) if numbers else None


def parse_simple_vertical_sheet(ws):
    products = []
    for row in ws.iter_rows(values_only=False):
        for cell in row:
            name = clean(cell.value)
            if cell.column != 2 or not name or is_label(name) or re.match(r"^\d", name) or len(name) < 3:
                continue
            price = None
            spec = ""
            for rr in range(cell.row + 1, min(ws.max_row, cell.row + 10) + 1):
                value = clean(ws.cell(rr, cell.column).value)
                if value and not spec and not is_label(value) and not re.search(r"\d+\s*-", value):
                    spec = value
                if value and ("PCS" in value.upper() or "P" in value.upper()):
                    price = first_price_from_line(value)
                    if price:
                        break
            if price:
                products.append(
                    {
                        "name": name,
                        "cost": price,
                        "colors": [],
                        "spec": spec,
                        "row": cell.row,
                        "col": cell.column,
                    }
                )
    return products


def parse_grid_sheet(ws):
    headers = []
    for row in ws.iter_rows(values_only=False):
        for cell in row:
            if clean(cell.value).upper() == "ITEM NAME":
                headers.append((cell.row, cell.column))
    if not headers:
        return parse_simple_vertical_sheet(ws)

    products = []
    for header_row, item_col in headers:
        name = ""
        price = None
        spec = ""
        for rr in range(header_row + 1, min(ws.max_row, header_row + 18) + 1):
            value = clean(ws.cell(rr, item_col).value)
            if not value:
                continue
            if not name and not is_label(value) and not re.fullmatch(r"\d+(?:\.\d+)?", value):
                name = value
            if value.upper() == "SPECIFICATION":
                spec = clean(ws.cell(rr + 1, item_col).value)
            if "PRICE" in value.upper():
                for pr in range(rr + 1, min(ws.max_row, rr + 5) + 1):
                    raw = ws.cell(pr, item_col).value
                    if isinstance(raw, (int, float)):
                        price = float(raw)
                        break
                    raw_text = clean(raw)
                    if re.fullmatch(r"\d+(?:\.\d+)?", raw_text):
                        price = float(raw_text)
                        break
                break
        if not name or not price:
            continue

        colors = []
        color_col = item_col + 1
        for rr in range(header_row + 1, min(ws.max_row, header_row + 40) + 1):
            color = clean(ws.cell(rr, color_col).value)
            if color and not is_bad_color(color) and not re.fullmatch(r"\d+(?:\.\d+)?", color):
                colors.append(color)
        seen = set()
        unique_colors = []
        for color in colors:
            key = color.upper()
            if key not in seen:
                unique_colors.append(color)
                seen.add(key)
        products.append(
            {
                "name": name,
                "cost": price,
                "colors": unique_colors,
                "spec": spec,
                "row": header_row,
                "col": item_col,
            }
        )
    return products


def detect_brand(filename, sheet_title):
    text = f"{filename} {sheet_title}".upper()
    for brand in ["GILLE", "ZEBRA", "HNJ", "MOB", "MOTOBOX"]:
        if brand in text:
            return brand
    return "Luis One"


def category_for(filename, sheet_title):
    text = f"{filename} {sheet_title}".upper()
    if "ACCESSORIES" in text:
        return ("Motorcycle Parts", "Motorcycle Accessories", "Brackets")
    if "MOTOBOX" in text or "MOTORBOX" in text:
        return ("Motorcycle Parts", "Motorcycle Storage", "Top Boxes")
    if "KIDS" in text:
        return ("Motorcycle Parts", "Helmets", "Kids Helmets")
    if "CAP TYPE" in text:
        return ("Motorcycle Parts", "Helmets", "Cap Type Helmets")
    if "MOTOCROSS" in text:
        return ("Motorcycle Parts", "Helmets", "Motocross Helmets")
    if "DUAL" in text:
        return ("Motorcycle Parts", "Helmets", "Dual Sport Helmets")
    if "MODULAR" in text:
        return ("Motorcycle Parts", "Helmets", "Modular Helmets")
    if "FULL FACE" in text:
        return ("Motorcycle Parts", "Helmets", "Full Face Helmets")
    if "HALF FACE" in text or "HAIF FACE" in text:
        return ("Motorcycle Parts", "Helmets", "Half Face Helmets")
    return ("Motorcycle Parts", "Helmets", "Helmets")


def price_tiers(cost):
    retail = math.ceil(cost * 1.15)
    wholesale_6 = math.ceil(cost * 1.10)
    wholesale_13 = math.ceil(cost * 1.08)
    return retail, [
        {"min_qty": 1, "max_qty": 5, "unit_price": retail},
        {"min_qty": 6, "max_qty": 12, "unit_price": wholesale_6},
        {"min_qty": 13, "max_qty": None, "unit_price": wholesale_13},
    ]


def uses_srp_discount(brand, filename):
    upper_brand = brand.upper()
    upper_file = filename.upper()
    return upper_brand == "GILLE" or "SRP" in upper_file


def supplier_cost_from_price(brand, filename, price):
    if uses_srp_discount(brand, filename):
        return price * 0.70
    return price


CATEGORY_CACHE = {}


def get_or_create_category(name, parent_id=None, level=1, sort_order=0):
    slug = slugify(name)
    if slug in CATEGORY_CACHE:
        return CATEGORY_CACHE[slug]
    query = "?slug=eq." + urllib.parse.quote(slug) + "&select=*"
    existing = request_json("GET", "categories", query, prefer=None)
    if existing:
        CATEGORY_CACHE[slug] = existing[0]
        return existing[0]
    payload = {
        "name_en": name,
        "name_zh": None,
        "slug": slug,
        "parent_id": parent_id,
        "level": level,
        "active": True,
        "show_on_homepage": False,
        "show_in_navigation": level <= 2,
        "sort_order": sort_order,
        "template_type": "marketplace",
        "description": f"{name} products.",
    }
    created = request_json("POST", "categories", payload=payload)
    CATEGORY_CACHE[slug] = created[0]
    return created[0]


def ensure_category_path(path):
    parent = None
    result = []
    for index, name in enumerate(path):
        category = get_or_create_category(name, parent["id"] if parent else None, index + 1, 80 + index)
        result.append(category)
        parent = category
    return result


def upsert_product(record):
    sku = record["sku"]
    query = "?on_conflict=sku&select=*"
    payload = {
        "sku": sku,
        "name": record["name"],
        "slug": record["slug"],
        "category_id": record["category_id"],
        "subcategory_id": record["subcategory_id"],
        "child_category_id": record["child_category_id"],
        "brand": record["brand"],
        "model": record["model"],
        "moq": 1,
        "stock_status": "ready_stock",
        "lead_time": "To be confirmed",
        "image_url": None,
        "description": record["description"],
        "supplier_notes": None,
        "internal_cost_notes": record["internal_cost_notes"],
        "admin_notes": f"Imported from {record['source_file']} / {record['source_sheet']}.",
        "active": True,
        "retail_price": record["retail_price"],
    }
    return request_json("POST", "products", query, [payload], prefer="resolution=merge-duplicates,return=representation")[0]


def patch_product_image(product_id, image_url, improved_name=None):
    payload = {"image_url": image_url}
    if improved_name:
        payload["name"] = improved_name
    query = "?id=eq." + urllib.parse.quote(product_id)
    request_json("PATCH", "products", query, payload=payload)


def replace_product_tiers(product_id, tiers):
    query = "?product_id=eq." + urllib.parse.quote(product_id)
    request_json("DELETE", "product_price_tiers", query, prefer=None)
    payload = [{"product_id": product_id, **tier} for tier in tiers]
    request_json("POST", "product_price_tiers", payload=payload)


def upsert_variant(product_id, variant_name, variant_sku, sort_order):
    payload = {
        "product_id": product_id,
        "variant_name": variant_name,
        "variant_sku": variant_sku,
        "model": variant_name,
        "fits": None,
        "image_url": None,
        "moq": 1,
        "stock_status": "ready_stock",
        "lead_time": "To be confirmed",
        "active": True,
        "sort_order": sort_order,
    }
    return request_json(
        "POST",
        "product_variants",
        "?on_conflict=variant_sku&select=*",
        [payload],
        prefer="resolution=merge-duplicates,return=representation",
    )[0]


def replace_variant_tiers(variant_id, tiers):
    query = "?variant_id=eq." + urllib.parse.quote(variant_id)
    request_json("DELETE", "product_variant_price_tiers", query, prefer=None)
    payload = [{"variant_id": variant_id, **tier} for tier in tiers]
    request_json("POST", "product_variant_price_tiers", payload=payload)


def patch_variants_image(product_id, image_url):
    query = "?product_id=eq." + urllib.parse.quote(product_id)
    request_json("PATCH", "product_variants", query, payload={"image_url": image_url}, prefer=None)


DRAWING_NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}


def read_relationships(zip_file, rel_path):
    if rel_path not in zip_file.namelist():
        return {}
    root = ET.fromstring(zip_file.read(rel_path))
    return {el.attrib["Id"]: el.attrib["Target"] for el in root}


def sheet_drawing_paths(zip_file):
    result = {}
    workbook_rels = read_relationships(zip_file, "xl/_rels/workbook.xml.rels")
    workbook = ET.fromstring(zip_file.read("xl/workbook.xml"))
    wb_ns = {
        "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }
    sheet_targets = {}
    for sheet in workbook.findall("main:sheets/main:sheet", wb_ns):
        name = sheet.attrib["name"]
        rid = sheet.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
        target = workbook_rels.get(rid, "")
        sheet_targets[name] = posixpath.normpath(posixpath.join("xl", target))

    for name, sheet_path in sheet_targets.items():
        rel_path = posixpath.join(posixpath.dirname(sheet_path), "_rels", posixpath.basename(sheet_path) + ".rels")
        rels = read_relationships(zip_file, rel_path)
        for target in rels.values():
            if "drawing" in target:
                result[name] = posixpath.normpath(posixpath.join(posixpath.dirname(sheet_path), target))
                break
    return result


def extract_sheet_images(xlsx_path):
    by_sheet = {}
    with zipfile.ZipFile(xlsx_path) as zip_file:
        drawings = sheet_drawing_paths(zip_file)
        for sheet_name, drawing_path in drawings.items():
            rel_path = posixpath.join(
                posixpath.dirname(drawing_path),
                "_rels",
                posixpath.basename(drawing_path) + ".rels",
            )
            rels = read_relationships(zip_file, rel_path)
            root = ET.fromstring(zip_file.read(drawing_path))
            images = []
            for anchor in root.findall("xdr:twoCellAnchor", DRAWING_NS) + root.findall("xdr:oneCellAnchor", DRAWING_NS):
                frm = anchor.find("xdr:from", DRAWING_NS)
                if frm is None:
                    continue
                col = int(frm.find("xdr:col", DRAWING_NS).text) + 1
                row = int(frm.find("xdr:row", DRAWING_NS).text) + 1
                cnv = anchor.find(".//xdr:cNvPr", DRAWING_NS)
                descr = clean(cnv.attrib.get("descr", "")) if cnv is not None else ""
                blip = anchor.find(".//a:blip", DRAWING_NS)
                if blip is None:
                    continue
                rid = blip.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                target = rels.get(rid)
                if not target:
                    continue
                media_path = posixpath.normpath(posixpath.join(posixpath.dirname(drawing_path), target))
                if media_path.startswith("xl/drawings/"):
                    media_path = media_path.replace("xl/drawings/../", "xl/")
                if media_path not in zip_file.namelist():
                    continue
                ext = os.path.splitext(media_path)[1].lower() or ".jpg"
                images.append(
                    {
                        "row": row,
                        "col": col,
                        "descr": descr,
                        "ext": ".jpg" if ext == ".jpeg" else ext,
                        "bytes": zip_file.read(media_path),
                    }
                )
            by_sheet[sheet_name] = images
    return by_sheet


def find_image_for_record(record, images):
    if not images:
        return None
    product_col = record["source_col"]
    product_row = record["source_row"]
    candidates = []
    for image in images:
        col_distance = abs(image["col"] - product_col)
        row_penalty = 0 if image["row"] <= product_row else 20
        candidates.append((col_distance, row_penalty, abs(product_row - image["row"]), image))
    candidates.sort(key=lambda item: (item[0], item[1], item[2]))
    return candidates[0][3] if candidates else None


def upload_image(record, image):
    if not image:
        return None
    image_bytes = optimize_image_bytes(image["bytes"])
    ext = ".jpg"
    storage_path = f"products/imported-helmets/{record['slug']}/{stable_suffix(record['sku'], image['descr'], len(image['bytes']))}{ext}"
    url = f"{SUPABASE_URL}/storage/v1/object/product-images/{urllib.parse.quote(storage_path, safe='/')}"
    content_type = {
        ".png": "image/png",
        ".webp": "image/webp",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
    }.get(ext, "image/jpeg")
    headers = {
        "apikey": SERVICE_KEY,
        "Authorization": f"Bearer {SERVICE_KEY}",
        "Content-Type": content_type,
        "x-upsert": "true",
    }
    req = urllib.request.Request(url, data=image_bytes, headers=headers, method="POST")
    try:
        with urllib.request.urlopen(req, timeout=120):
            pass
    except urllib.error.HTTPError as exc:
        details = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"Storage upload failed: {exc.code} {details}") from exc
    return f"{SUPABASE_URL}/storage/v1/object/public/product-images/{storage_path}"


def optimize_image_bytes(raw_bytes):
    max_side = 1600
    max_bytes = 1_800_000
    try:
        source = Image.open(BytesIO(raw_bytes))
        source = source.convert("RGB")
        source.thumbnail((max_side, max_side), Image.Resampling.LANCZOS)
        for quality in (88, 82, 76, 70, 64):
            output = BytesIO()
            source.save(output, format="JPEG", quality=quality, optimize=True, progressive=True)
            data = output.getvalue()
            if len(data) <= max_bytes:
                return data
        return data
    except Exception:
        return raw_bytes


def improved_name(record, image):
    descr = clean((image or {}).get("descr", ""))
    if not descr:
        return record["name"]
    normalized = descr.replace("-", " ")
    normalized = re.sub(r"\s+", " ", normalized).strip()
    if not normalized or normalized.upper() in record["name"].upper():
        return record["name"]
    if record["name"].upper() in normalized.upper():
        return normalized
    return record["name"]


def collect_products():
    records = []
    for filename in FILES:
        path = os.path.join(BASE_DIR, filename)
        wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
        try:
            for ws in wb.worksheets:
                category_path = category_for(filename, ws.title)
                brand = detect_brand(filename, ws.title)
                for product in parse_grid_sheet(ws):
                    source_price = float(product["cost"])
                    cost = supplier_cost_from_price(brand, filename, source_price)
                    retail_price, tiers = price_tiers(cost)
                    suffix = stable_suffix(filename, ws.title, product["row"], product["col"], product["name"])
                    base_slug = slugify(product["name"])
                    sku = f"{brand}-{base_slug}-{suffix}".upper()[:64]
                    records.append(
                        {
                            "sku": sku,
                            "slug": f"{base_slug}-{suffix}",
                            "name": product["name"],
                            "brand": brand,
                            "model": product["name"],
                            "cost": cost,
                            "source_price": source_price,
                            "retail_price": retail_price,
                            "tiers": tiers,
                            "colors": product["colors"],
                            "description": (
                                f"{product['name']} from {brand}. Available colors are listed as selectable variants. "
                                "Wholesale order is manually confirmed by Luis One Supply Hub."
                            ),
                            "category_path": category_path,
                            "source_file": filename,
                            "source_sheet": ws.title,
                            "source_row": product["row"],
                            "source_col": product["col"],
                            "internal_cost_notes": (
                                f"Imported SRP: PHP {source_price:.2f}. Supplier discount 30%; "
                                f"calculated cost PHP {cost:.2f}. Retail +15%, wholesale 6-12 +10%, 13+ +8%."
                                if uses_srp_discount(brand, filename)
                                else f"Imported supplier cost: PHP {cost:.2f}. Retail +15%, wholesale 6-12 +10%, 13+ +8%."
                            ),
                        }
                    )
        finally:
            wb.close()
    return records


def main():
    dry_run = "--dry-run" in sys.argv
    records = collect_products()
    total_variants = sum(len(r["colors"]) for r in records)
    print(f"Parsed products: {len(records)}")
    print(f"Parsed variants/colors: {total_variants}")
    print("Sample:")
    for record in records[:8]:
        print(
            f"- {record['sku']} | {record['name']} | cost {record['cost']:.2f} | retail {record['retail_price']} | variants {len(record['colors'])}"
        )
    if dry_run:
        return

    created_or_updated = 0
    variants_saved = 0
    categories_seen = set()
    image_maps = {}
    if "--with-images" in sys.argv:
        for filename in FILES:
            path = os.path.join(BASE_DIR, filename)
            image_maps[filename] = extract_sheet_images(path)
            image_count = sum(len(images) for images in image_maps[filename].values())
            print(f"Images found in {filename}: {image_count}")

    images_uploaded = 0
    for index, record in enumerate(records, start=1):
        categories = ensure_category_path(record["category_path"])
        for category in categories:
            categories_seen.add(category["slug"])
        record["category_id"] = categories[0]["id"]
        record["subcategory_id"] = categories[1]["id"] if len(categories) > 1 else None
        record["child_category_id"] = categories[2]["id"] if len(categories) > 2 else None

        db_product = upsert_product(record)
        replace_product_tiers(db_product["id"], record["tiers"])
        created_or_updated += 1

        image_url = None
        if image_maps:
            images = image_maps.get(record["source_file"], {}).get(record["source_sheet"], [])
            image = find_image_for_record(record, images)
            image_url = upload_image(record, image) if image else None
            if image_url:
                patch_product_image(db_product["id"], image_url, improved_name(record, image))
                images_uploaded += 1

        for variant_index, color in enumerate(record["colors"], start=1):
            variant_slug = slugify(color, f"variant-{variant_index}")[:36]
            variant_sku = f"{record['sku']}-{variant_slug}".upper()[:96]
            db_variant = upsert_variant(db_product["id"], color, variant_sku, variant_index)
            replace_variant_tiers(db_variant["id"], record["tiers"])
            variants_saved += 1
        if image_url:
            patch_variants_image(db_product["id"], image_url)

        if index % 25 == 0:
            print(f"Imported {index}/{len(records)} products...")
            time.sleep(0.2)

    print(f"Done. Products created/updated: {created_or_updated}")
    print(f"Variants created/updated: {variants_saved}")
    print(f"Product images uploaded/updated: {images_uploaded}")
    print(f"Categories used/created: {len(categories_seen)}")


if __name__ == "__main__":
    main()
