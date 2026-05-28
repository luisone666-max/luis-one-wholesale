import importlib.util
import posixpath
import time
import urllib.request
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
BASE_SCRIPT = ROOT / "scripts" / "prepare-google-ok-preview.py"
DOWNLOAD_DIR = ROOT / "tmp-imports" / "google-ok" / "downloads"
PUBLIC_IMAGE_ROOT = ROOT / "public" / "product-images" / "google-ok"

spec = importlib.util.spec_from_file_location("google_ok_preview_base", BASE_SCRIPT)
base = importlib.util.module_from_spec(spec)
spec.loader.exec_module(base)


NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
}


def xlsx_path_for(gid, sheet_name):
    return DOWNLOAD_DIR / f"{gid}-{base.slugify(sheet_name)}.xlsx"


def download_sheet_xlsx(gid, sheet_name):
    DOWNLOAD_DIR.mkdir(parents=True, exist_ok=True)
    path = xlsx_path_for(gid, sheet_name)
    if path.exists() and path.stat().st_size > 1000:
        return path

    url = f"https://docs.google.com/spreadsheets/d/{base.SPREADSHEET_ID}/export?format=xlsx&gid={gid}"
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    temp_path = path.with_suffix(".xlsx.tmp")

    with urllib.request.urlopen(request, timeout=240) as response, temp_path.open("wb") as fh:
        while True:
            chunk = response.read(1024 * 1024)
            if not chunk:
                break
            fh.write(chunk)

    temp_path.replace(path)
    time.sleep(0.15)
    return path


def rels_for(zip_file, rels_path):
    if rels_path not in zip_file.namelist():
        return {}
    root = ET.fromstring(zip_file.read(rels_path))
    return {rel.attrib["Id"]: rel.attrib["Target"] for rel in root.findall("rel:Relationship", NS)}


def resolve_target(source_path, target):
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(source_path), target))


def worksheet_paths(zip_file):
    return sorted(
        name
        for name in zip_file.namelist()
        if name.startswith("xl/worksheets/sheet") and name.endswith(".xml") and "/_rels/" not in name
    )


def drawing_paths_for_workbook(zip_file):
    paths = []
    for worksheet_path in worksheet_paths(zip_file):
        worksheet_root = ET.fromstring(zip_file.read(worksheet_path))
        drawing = worksheet_root.find(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}drawing")
        if drawing is None:
            continue
        rel_id = drawing.attrib.get(f"{{{NS['r']}}}id")
        if not rel_id:
            continue
        rels_path = posixpath.join(posixpath.dirname(worksheet_path), "_rels", posixpath.basename(worksheet_path) + ".rels")
        target = rels_for(zip_file, rels_path).get(rel_id)
        if target:
            paths.append(resolve_target(worksheet_path, target))
    return paths


def extract_images_by_row(xlsx_path, sheet_name):
    image_dir = PUBLIC_IMAGE_ROOT / base.slugify(sheet_name)
    image_dir.mkdir(parents=True, exist_ok=True)
    images_by_row = {}

    with zipfile.ZipFile(xlsx_path) as zip_file:
        for drawing_path in drawing_paths_for_workbook(zip_file):
            if drawing_path not in zip_file.namelist():
                continue
            rels_path = posixpath.join(posixpath.dirname(drawing_path), "_rels", posixpath.basename(drawing_path) + ".rels")
            drawing_rels = rels_for(zip_file, rels_path)
            drawing_root = ET.fromstring(zip_file.read(drawing_path))

            for anchor in list(drawing_root):
                marker = anchor.find("xdr:from", NS)
                blip = anchor.find(".//a:blip", NS)
                if marker is None or blip is None:
                    continue
                rel_id = blip.attrib.get(f"{{{NS['r']}}}embed")
                target = drawing_rels.get(rel_id)
                if not target:
                    continue

                media_path = resolve_target(drawing_path, target)
                if media_path not in zip_file.namelist():
                    continue

                row = int(marker.find("xdr:row", NS).text) + 1
                col = int(marker.find("xdr:col", NS).text) + 1
                media_name = Path(media_path).name
                output_name = f"r{row:04d}-c{col:02d}-{media_name}".lower()
                output_path = image_dir / output_name
                data = zip_file.read(media_path)
                if not output_path.exists() or output_path.stat().st_size != len(data):
                    output_path.write_bytes(data)
                images_by_row.setdefault(row, []).append(f"/product-images/google-ok/{base.slugify(sheet_name)}/{output_name}")

    return images_by_row


def workbook_rows(xlsx_path):
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True)
    worksheet = workbook.active
    rows = []
    for row_index in range(1, worksheet.max_row + 1):
        cells = []
        for col_index in range(1, worksheet.max_column + 1):
            cells.append({"text": base.clean(worksheet.cell(row_index, col_index).value), "images": []})
        rows.append({"row": row_index, "cells": cells})
    return rows


def attach_row_images(rows, images_by_row):
    for row in rows:
        images = images_by_row.get(row["row"], [])
        if images and row["cells"]:
            row["cells"][0]["images"] = images
    return rows


def parse_products_from_sheet(gid, sheet_name):
    xlsx_path = download_sheet_xlsx(gid, sheet_name)
    rows = attach_row_images(workbook_rows(xlsx_path), extract_images_by_row(xlsx_path, sheet_name))

    images_by_code = {}
    for row in rows:
        code, _ = base.first_code_and_price(row["cells"])
        images = base.cell_images(row["cells"])
        if code and images:
            images_by_code[code] = images

    active_header = {}
    current_title = base.sheet_display_name(sheet_name)
    last_image_for_title = {}
    variants = []

    for row in rows:
        cells = row["cells"]
        texts = [cell["text"] for cell in cells]
        maybe_header = base.header_map(cells)
        if maybe_header:
            active_header = maybe_header
            continue

        if base.is_title_row(texts):
            values = [base.clean(text) for text in texts if base.clean(text)]
            title = " ".join(values[:2])
            if title.upper() not in base.NON_PRODUCT_TITLES:
                current_title = title
            continue

        if not active_header:
            continue

        code = base.cell_text(cells, active_header.get("code"))
        qty_box, price = base.infer_qty_and_price(
            cells,
            active_header.get("code", 0),
            active_header.get("qty"),
            active_header.get("price"),
        )
        if not base.looks_like_code(code) or price is None:
            continue

        attrs = {}
        for index, label in active_header.get("attrs", []):
            value = base.cell_text(cells, index)
            if value:
                attrs[label.lower().replace("/", "_")] = value

        images = images_by_code.get(code, base.cell_images(cells))
        title_key = f"{sheet_name}|{current_title}"
        if images:
            last_image_for_title[title_key] = images[0]
        image_url = images[0] if images else last_image_for_title.get(title_key, "")
        variants.append(
            {
                "sheet_id": gid,
                "sheet": sheet_name,
                "row": row["row"],
                "source_title": current_title,
                "code": code,
                "color": attrs.get("color") or attrs.get("colour") or "",
                "model": attrs.get("model")
                or attrs.get("fitment")
                or attrs.get("fit")
                or attrs.get("description")
                or attrs.get("brand")
                or "",
                "size": attrs.get("size") or attrs.get("type") or attrs.get("pcs") or attrs.get("specification") or "",
                "qty_box": qty_box,
                "cost": price,
                "source_image_url": image_url,
            }
        )

    grouped = {}
    for variant in variants:
        key = (sheet_name, variant["source_title"])
        grouped.setdefault(key, []).append(variant)
    return grouped


def main():
    base.parse_products_from_sheet = parse_products_from_sheet
    base.main()


if __name__ == "__main__":
    main()
